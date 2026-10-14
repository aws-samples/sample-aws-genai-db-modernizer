"""Key figures of every deliverable, read back from what each one renders (#573).

AGENTS.md "Deliverables agree": the decision report, engineering report,
deck/PDF and the interactive analysis report all render one synthesis
``report.json``, so a figure shown in two of them must be the same figure.
This module reads the key figures back out of each rendered deliverable and
computes what ``report.json`` says they should be, so
``test_deliverable_figures.py`` can compare them one by one.

Extraction reads stable hooks, never free prose: ``data-*`` attributes in the
two HTML reports (``data-wave``, ``data-figure``, ``data-engine``), and fixed
labels in the Markdown, deck and PDF ("Projected monthly (target):",
"N source tables migrate", "Kept in sync from X: N tables", ...). The one
deliberate exception is the summary's "Other targets evaluated:" clause, the
only place three of the deliverables name an engine the reality check
removed; it is a fixed label too.

A figure is keyed by a short string (``"cost:dynamodb"``, ``"wave3:moving"``,
``"risks:HIGH"``...). Each deliverable records every observation of a key it
shows, with where it was seen, so one key shown twice in one deliverable is
checked twice.
"""

from __future__ import annotations

import html as html_lib
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Any

from src.shared.engine_names import ENGINE_DISPLAY_NAMES

SEVERITIES = ("CRITICAL", "HIGH", "MEDIUM", "LOW")

# Recorded when a hook or label is there but its figure is not, so the
# comparison reports "is '<unreadable>'" instead of passing silently.
UNREADABLE = "<unreadable>"

# Display name or raw id -> engine id. Longest first, so "Aurora PostgreSQL"
# wins over "Aurora" in an alternation.
_NAME_TO_ENGINE: dict[str, str] = {}
for _eid, _name in ENGINE_DISPLAY_NAMES.items():
    _NAME_TO_ENGINE[_name] = _eid
    _NAME_TO_ENGINE[_eid] = _eid
_NAMES_RE = "|".join(re.escape(n) for n in sorted(_NAME_TO_ENGINE, key=len, reverse=True))


def engine_id(name: str) -> str:
    """Engine id for a display name or raw id; the input unchanged if unknown."""
    return _NAME_TO_ENGINE.get(name.strip(), name.strip())


def engines_in(text: str) -> set[str]:
    """Every engine named (by display name or raw id) in ``text``."""
    return {engine_id(m) for m in re.findall(rf"(?<![\w-])(?:{_NAMES_RE})(?![\w-])", text)}


def norm(text: str) -> str:
    """Collapse whitespace runs: the PDF text layer and deck shapes break lines
    differently, the figures do not."""
    return " ".join(text.split())


def _int(s: str) -> int:
    return int(s.replace(",", ""))


def _usd(s: str) -> float:
    return round(float(s.replace("$", "").replace(",", "")), 2)


@dataclass
class Figures:
    """The figures one deliverable shows. ``values[key]`` is every
    ``(value, where)`` observation of ``key``; ``engines`` maps every engine
    the deliverable names in a structural slot (an engine row, a cost card, a
    wave's engines, a risk's engine, a tied query's engines) to where."""

    deliverable: str
    values: dict[str, list[tuple[Any, str]]] = field(default_factory=dict)
    engines: dict[str, set[str]] = field(default_factory=dict)

    def add(self, key: str, value: Any, where: str) -> None:
        self.values.setdefault(key, []).append((value, where))

    def name(self, engines: set[str] | list[str] | str, where: str) -> None:
        for e in [engines] if isinstance(engines, str) else engines:
            self.engines.setdefault(e, set()).add(where)


# --------------------------------------------------------------------------
# Expected figures, straight from report.json
# --------------------------------------------------------------------------


def final_architecture(report: dict[str, Any]) -> set[str]:
    """Engines in the final architecture: every engine that owns queries after
    the reality check, plus the cache layer (it owns none, #296)."""
    after = (report.get("reality_check") or {}).get("after_distribution") or {}
    engines = {e for e, n in after.items() if n}
    cache = (report.get("cache_overlay") or {}).get("engine")
    if cache:
        engines.add(cache)
    return engines


def removed_engines(report: dict[str, Any]) -> set[str]:
    """Engines the reality check removed: they owned queries before it, or a
    full consolidation moved them out, and none of them is in the final
    architecture."""
    rc = report.get("reality_check") or {}
    before = {e for e, n in (rc.get("before_distribution") or {}).items() if n}
    consolidated = {
        c.get("from_engine")
        for c in rc.get("consolidations") or []
        if isinstance(c, dict) and c.get("action") == "full"
    }
    return {e for e in before | consolidated if e} - final_architecture(report)


def migrating(wave: dict[str, Any]) -> bool:
    """A wave that moves tables off Aurora to a purpose-built engine. Wave 1
    (``homogeneity`` set) carries the whole schema to Aurora and is not one."""
    return bool(wave.get("moves_from")) and wave.get("homogeneity") is None


def wave_synced_tables(wave: dict[str, Any]) -> set[str]:
    """Tables a wave keeps in sync from an owner outside it. In the merged
    DynamoDB/DocumentDB wave, a table one engine serves from the other moves
    with its owner (#581): it is not kept in sync from anyone outside."""
    engines = set(wave.get("engines") or [])
    return {
        str(o.get("table"))
        for o in wave.get("table_owners") or []
        if isinstance(o, dict) and o.get("owner") not in engines
    }


def wave_moving_by_engine(wave: dict[str, Any]) -> dict[str, int]:
    """Tables each engine of a migrating wave owns, i.e. moves."""
    owned = wave.get("owned_counts_by_engine")
    if owned:
        return {e: int(n) for e, n in owned.items()}
    engines = wave.get("engines") or []
    if len(engines) == 1:
        return {engines[0]: int(wave.get("table_count") or 0) - len(wave.get("table_owners") or [])}
    return {}


def wave_served_by_engine(wave: dict[str, Any]) -> dict[str, int]:
    """Tables each engine of a migrating wave serves."""
    served = wave.get("table_counts_by_engine")
    if served:
        return {e: int(n) for e, n in served.items()}
    engines = wave.get("engines") or []
    if len(engines) == 1:
        return {engines[0]: int(wave.get("table_count") or 0)}
    return {}


def expected_figures(report: dict[str, Any]) -> dict[str, Any]:
    """Every key figure, computed from report.json fields only (never through a
    renderer helper, so a renderer bug cannot hide in both sides)."""
    exp: dict[str, Any] = {}
    after = (report.get("reality_check") or {}).get("after_distribution") or {}
    for e, n in after.items():
        if n:
            exp[f"queries:{e}"] = int(n)
    overlay = report.get("cache_overlay") or {}
    if overlay.get("engine"):
        exp[f"cached_reads:{overlay['engine']}"] = int(overlay.get("query_count") or 0)

    tables_migrate = 0
    for w in report.get("migration_waves") or []:
        n = w.get("wave")
        exp[f"wave{n}:queries"] = int(w.get("query_count") or 0)
        if w.get("cutover_query_count") is not None:
            exp[f"wave{n}:cutover_queries"] = int(w["cutover_query_count"])
        served = int(w.get("table_count") or 0)
        if served:
            exp[f"wave{n}:served"] = served
        if w.get("homogeneity") is not None and report.get("source_table_count") is not None:
            # #578: wave 1 carries every analyzed table plus the real views.
            exp[f"wave{n}:tables_only"] = int(report["source_table_count"])
            exp[f"wave{n}:views"] = int(((report.get("facts") or {}).get("source_view_count")) or 0)
        if migrating(w) and served:
            per_engine = wave_moving_by_engine(w)
            # #581: the tables that move are the ones the wave's engines own.
            moving = sum(per_engine.values()) if per_engine else served - len(wave_synced_tables(w))
            exp[f"wave{n}:synced"] = served - moving
            exp[f"wave{n}:moving"] = moving
            tables_migrate += moving
            for e, k in wave_moving_by_engine(w).items():
                exp[f"wave{n}:moving:{e}"] = k
            for e, k in wave_served_by_engine(w).items():
                exp[f"tables:{e}"] = exp.get(f"tables:{e}", 0) + k
    exp["tables_migrate"] = tables_migrate

    tco = report.get("tco_analysis") or {}
    for row in tco.get("cost_breakdown") or []:
        if isinstance(row, dict) and isinstance(row.get("monthly_cost_usd"), (int, float)):
            exp[f"cost:{row['database']}"] = round(float(row["monthly_cost_usd"]), 2)
    if isinstance(tco.get("projected_monthly_cost"), (int, float)):
        exp["cost:total"] = round(float(tco["projected_monthly_cost"]), 2)

    risks = [r for r in (report.get("risk_assessment") or {}).get("risks") or [] if r]
    counts = Counter(str(r.get("severity", "")).upper() for r in risks)
    for sev in SEVERITIES:
        exp[f"risks:{sev}"] = counts.get(sev, 0)
    exp["risks:total"] = len(risks)

    exp["tied"] = len([t for t in report.get("tied_queries") or [] if isinstance(t, dict)])
    exp["engine_count"] = len(final_architecture(report))
    exp["removed"] = frozenset(removed_engines(report))
    # A pointer to the decision report's "Engines not used" section must have
    # the section to point at (#547).
    exp["engines_not_used_section"] = True
    return exp


# --------------------------------------------------------------------------
# Invariants: no deliverable may show an impossible figure
# --------------------------------------------------------------------------


def invariant_violations(
    report: dict[str, Any], figs: dict[str, list[tuple[Any, str]]], label: str
) -> list[tuple[str, str]]:
    """``(key, message)`` for every figure in ``figs`` that breaks an invariant:
    a table total above the tables analyzed, a query total above the queries
    analyzed, a wave whose moved and kept-in-sync tables do not add up to the
    tables it serves."""
    tables_analyzed = int(report.get("source_table_count") or 0)
    views = int(((report.get("facts") or {}).get("source_view_count")) or 0)
    whole_schema = {
        f"wave{w.get('wave')}:served"
        for w in report.get("migration_waves") or []
        if isinstance(w, dict) and w.get("homogeneity") is not None
    }
    queries_analyzed = int((report.get("assignment_summary") or {}).get("query_count") or 0)
    out: list[tuple[str, str]] = []
    for key, observations in figs.items():
        is_table = bool(re.match(r"tables|wave\d+:(served|moving|synced)", key))
        is_query = bool(
            re.match(
                r"(queries|cached_reads):|tied$|wave\d+:(queries|cutover_queries|group_queries)$",
                key,
            )
        )
        for value, where in observations:
            if not isinstance(value, int) or isinstance(value, bool):
                continue
            # Wave 1 carries the whole schema: every table and view (#578).
            limit = tables_analyzed + (views if key in whole_schema else 0)
            if is_table and tables_analyzed and value > limit:
                out.append(
                    (
                        key,
                        f"{label} {where}: {key}={value} exceeds the {limit} "
                        "source tables (and views, for wave 1) analyzed",
                    )
                )
            if is_query and queries_analyzed and value > queries_analyzed:
                out.append(
                    (
                        key,
                        f"{label} {where}: {key}={value} exceeds the {queries_analyzed} "
                        "queries analyzed",
                    )
                )
        m = re.match(r"(wave\d+):group_queries$", key)
        if m:
            for grouped, where in observations:
                for total, _ in figs.get(f"{m.group(1)}:queries", []):
                    if grouped > total:
                        out.append(
                            (
                                key,
                                f"{label} {where}: its table groups hold {grouped} queries, "
                                f"more than the wave's {total}",
                            )
                        )
    for key in figs:
        m = re.match(r"(wave\d+):served$", key)
        if not m:
            continue
        w = m.group(1)
        for served, where in figs[key]:
            for moving, _ in figs.get(f"{w}:moving", []):
                for synced, _ in figs.get(f"{w}:synced", []):
                    if moving + synced != served:
                        out.append(
                            (
                                f"{w}:reconcile",
                                f"{label} {where}: {w} moves {moving} and keeps {synced} in "
                                f"sync, which is not the {served} it serves",
                            )
                        )
    return out


def report_invariant_violations(report: dict[str, Any]) -> list[tuple[str, str]]:
    """Invariants on report.json itself: wave counts reconcile with their
    table groups, sync lists and per-engine splits; costs add up; every engine
    it names is in the final architecture."""
    out: list[tuple[str, str]] = []
    final = final_architecture(report)
    queries_analyzed = int((report.get("assignment_summary") or {}).get("query_count") or 0)

    ranking = {str(r.get("target")) for r in report.get("ranking") or [] if isinstance(r, dict)}
    if ranking != final:
        out.append(("ranking", f"ranking engines {sorted(ranking)} != final {sorted(final)}"))
    for r in report.get("ranking") or []:
        after = (report.get("reality_check") or {}).get("after_distribution") or {}
        if r.get("target") in after and r.get("assigned_queries") != after[r["target"]]:
            out.append(
                (
                    f"queries:{r['target']}",
                    f"ranking says {r['target']} owns {r.get('assigned_queries')} queries, "
                    f"after_distribution says {after[r['target']]}",
                )
            )
    owned_total = sum(
        ((report.get("reality_check") or {}).get("after_distribution") or {}).values()
    )
    if queries_analyzed and owned_total != queries_analyzed:
        out.append(
            (
                "queries:total",
                f"owned queries add up to {owned_total}, not the {queries_analyzed} analyzed",
            )
        )

    for w in report.get("migration_waves") or []:
        n = w.get("wave")
        tables = list(w.get("tables") or [])
        if len(set(tables)) != len(tables):
            out.append((f"wave{n}:tables", f"wave {n} lists a table more than once"))
        if w.get("table_count") is not None and int(w["table_count"]) != len(tables):
            out.append(
                (
                    f"wave{n}:served",
                    f"wave {n} table_count {w['table_count']} != {len(tables)} tables listed",
                )
            )
        unknown = set(w.get("engines") or []) - final
        if unknown:
            out.append(
                (
                    f"wave{n}:engines",
                    f"wave {n} names {sorted(unknown)}, not in the final architecture",
                )
            )
        group_tables: list[str] = []
        group_queries = 0
        for g in w.get("table_groups") or []:
            group_tables += list(g.get("tables") or [])
            group_queries += int(g.get("query_count") or 0)
        if len(set(group_tables)) != len(group_tables):
            out.append((f"wave{n}:groups", f"wave {n} puts a table in two table groups"))
        stray = set(group_tables) - set(tables)
        if stray:
            out.append(
                (
                    f"wave{n}:groups",
                    f"wave {n} table groups hold {len(stray)} tables the wave does not list",
                )
            )
        if group_queries > int(w.get("query_count") or 0):
            out.append(
                (
                    f"wave{n}:groups",
                    f"wave {n} table groups hold {group_queries} queries, more than its "
                    f"{w.get('query_count')}",
                )
            )
        if w.get("homogeneity") is not None and report.get("source_table_count") is not None:
            views = int(((report.get("facts") or {}).get("source_view_count")) or 0)
            if int(w.get("table_count") or 0) != int(report["source_table_count"]) + views:
                out.append(
                    (
                        f"wave{n}:served",
                        f"wave {n} carries {w.get('table_count')} tables and views, not the "
                        f"{report['source_table_count']} tables analyzed plus {views} views",
                    )
                )
        owners = w.get("table_owners") or []
        owner_tables = [o.get("table") for o in owners]
        if set(owner_tables) - set(tables):
            out.append((f"wave{n}:synced", f"wave {n} keeps in sync tables it does not serve"))
        # A single-engine wave never lists its own engine as a table's owner;
        # the merged wave may (one engine serves what the other owns, #581).
        if len(w.get("engines") or []) == 1 and {o.get("owner") for o in owners} & set(
            w.get("engines") or []
        ):
            out.append((f"wave{n}:synced", f"wave {n} keeps a table in sync from its own engine"))
        if {o.get("owner") for o in owners} - final:
            out.append(
                (f"wave{n}:synced", f"wave {n} syncs from an engine not in the architecture")
            )
        if migrating(w) and w.get("table_count"):
            synced = wave_synced_tables(w)
            moving = int(w["table_count"]) - len(synced)
            per_engine = wave_moving_by_engine(w)
            if per_engine and sum(per_engine.values()) != moving:
                out.append(
                    (
                        f"wave{n}:moving",
                        f"wave {n} engines move {sum(per_engine.values())} tables between them, "
                        f"but it serves {w['table_count']} and keeps {len(synced)} in sync "
                        f"from outside it ({moving} move)",
                    )
                )
            served = wave_served_by_engine(w)
            for e, k in per_engine.items():
                if k > served.get(e, k):
                    out.append(
                        (f"wave{n}:moving:{e}", f"wave {n}: {e} moves {k} tables, serves fewer")
                    )

    tco = report.get("tco_analysis") or {}
    breakdown = {
        r["database"]: float(r["monthly_cost_usd"])
        for r in tco.get("cost_breakdown") or []
        if isinstance(r, dict) and isinstance(r.get("monthly_cost_usd"), (int, float))
    }
    if set(breakdown) - final:
        out.append(
            (
                "cost",
                f"cost breakdown prices {sorted(set(breakdown) - final)}, not in the architecture",
            )
        )
    if final - set(breakdown):
        out.append(("cost", f"cost breakdown has no price for {sorted(final - set(breakdown))}"))
    projected = tco.get("projected_monthly_cost")
    if isinstance(projected, (int, float)) and abs(sum(breakdown.values()) - projected) > 0.011:
        out.append(
            (
                "cost:total",
                f"per-engine costs add up to {sum(breakdown.values()):.2f}, not {projected}",
            )
        )

    for t in report.get("tied_queries") or []:
        named = set(t.get("tied_engines") or []) | {t.get("assigned_engine")}
        if named - final:
            out.append(("tied", f"tied query {t.get('query_id')} names {sorted(named - final)}"))
    summary_ties = (report.get("assignment_summary") or {}).get("tie_unresolved_count")
    if summary_ties is not None and summary_ties != len(report.get("tied_queries") or []):
        out.append(
            (
                "tied",
                f"tie_unresolved_count {summary_ties} != {len(report.get('tied_queries') or [])} tied queries",
            )
        )
    for r in (report.get("risk_assessment") or {}).get("risks") or []:
        m = re.match(r"\[(\w+)\]", str(r.get("description") or ""))
        if m and m.group(1) in ENGINE_DISPLAY_NAMES and m.group(1) not in final:
            out.append(("risks", f"{r.get('risk_id')} is on {m.group(1)}, not in the architecture"))
    return out


# --------------------------------------------------------------------------
# Decision report (HTML, read in a browser)
# --------------------------------------------------------------------------

_DECISION_WAVE_CUTOVER = re.compile(
    r"^all ([\d,]+) quer(?:y|ies) at cutover; ([\d,]+) \([\d.]+%\) remain on .+? at the end"
    r"(?:, ([\d,]+) source tables? and views(?: \(([\d,]+) tables?, ([\d,]+) views?\))?)?$"
)
_DECISION_WAVE = re.compile(
    r"^([\d,]+) quer(?:y|ies) \([\d.]+% of (?:calls|query patterns)\)"
    r"(?:, ([\d,]+) served, ([\d,]+) move|, ([\d,]+) source tables?)?$"
)
_OTHER_TARGETS = re.compile(r"Other targets evaluated: (.+?)\.(?:\s|$)")


def _other_targets(summary: str) -> set[str]:
    """Engines the summary's "Other targets evaluated: X (consolidated into Y
    ...), ..." clause names as evaluated and not kept."""
    m = _OTHER_TARGETS.search(summary)
    if not m:
        return set()
    out = set()
    # Drop the parenthesised "(consolidated into Y ...)" asides, then split.
    for name in re.sub(r"\([^)]*\)", "", m.group(1)).split(","):
        if name.strip():
            out.add(engine_id(name))
    return out


def decision_figures(page: Any) -> Figures:
    """Figures the decision report shows, read from the rendered page."""
    f = Figures("decision report")

    def tile(label: str) -> str:
        return str(
            page.locator(".tile", has=page.locator(f"p:text-is('{label}')"))
            .locator(".val")
            .inner_text()
        )

    f.add("cost:total", _usd(tile("Projected monthly")), "Projected monthly tile")
    f.add("engine_count", _int(tile("Engines")), "Engines tile")
    f.add("tables_migrate", _int(tile("Tables migrate")), "Tables migrate tile")

    for row in page.locator("tr[data-engine]").all():
        eng = row.get_attribute("data-engine")
        f.name(eng, "engine table")
        scope = norm(row.locator("[data-figure=scope]").inner_text())
        m = re.match(r"^([\d,]+) source tables?\b", scope)
        if m:
            f.add(f"tables:{eng}", _int(m.group(1)), "engine table Scope")
        cost = norm(row.locator("[data-figure=monthly-cost]").inner_text())
        if cost.startswith("$"):
            f.add(f"cost:{eng}", _usd(cost), "engine table Est. monthly")
        m = re.match(
            r"^([\d,]+) cached reads?\b", norm(row.locator("[data-figure=workload]").inner_text())
        )
        if m:
            f.add(f"cached_reads:{eng}", _int(m.group(1)), "engine table Workload")
    total = page.locator("tr[data-figure=total]")
    if total.count():
        f.add(
            "cost:total",
            _usd(norm(total.locator("[data-figure=monthly-cost]").inner_text())),
            "engine table Total",
        )
        m = re.match(
            r"^([\d,]+) tables? migrates?$", norm(total.locator("[data-figure=scope]").inner_text())
        )
        if m:
            f.add("tables_migrate", _int(m.group(1)), "engine table Total")

    for wave in page.locator("[data-wave]").all():
        n = wave.get_attribute("data-wave")
        stats = norm(wave.locator("[data-figure=wave-stats]").inner_text())
        where = f"roadmap wave {n}"
        f.name(engines_in(wave.locator("b").first.inner_text()), f"{where} title")
        m = _DECISION_WAVE_CUTOVER.match(stats)
        if m:
            f.add(f"wave{n}:cutover_queries", _int(m.group(1)), where)
            f.add(f"wave{n}:queries", _int(m.group(2)), where)
            if m.group(3):
                f.add(f"wave{n}:served", _int(m.group(3)), where)
            if m.group(4):
                f.add(f"wave{n}:tables_only", _int(m.group(4)), where)
                f.add(f"wave{n}:views", _int(m.group(5)), where)
            continue
        m = _DECISION_WAVE.match(stats)
        if not m:
            f.add(f"wave{n}:unparsed", stats, where)
            continue
        f.add(f"wave{n}:queries", _int(m.group(1)), where)
        if m.group(2):
            served, moving = _int(m.group(2)), _int(m.group(3))
            f.add(f"wave{n}:served", served, where)
            f.add(f"wave{n}:moving", moving, where)
            f.add(f"wave{n}:synced", served - moving, where)
        elif m.group(4):
            f.add(f"wave{n}:served", _int(m.group(4)), where)

    risk = page.locator("[data-figure=risk-summary]")
    if risk.count():
        text = norm(risk.inner_text())
        m = re.search(r"([\d,]+) open migration risks? \(([^)]*)\)", text)
        if m:
            f.add("risks:total", _int(m.group(1)), "Risk posture")
            sev = {s.split()[1].upper(): _int(s.split()[0]) for s in m.group(2).split(", ")}
            for s in SEVERITIES:
                f.add(f"risks:{s}", sev.get(s, 0), "Risk posture")

    tied = page.locator("[data-figure=tied-queries]")
    if tied.count():
        m = re.search(r"\(([\d,]+) quer(?:y|ies)\)", tied.locator("h2").inner_text())
        f.add("tied", _int(m.group(1)) if m else UNREADABLE, "Pending engine choices heading")
        items = tied.locator("li")
        f.add("tied", items.count(), "Pending engine choices list")
        for li in items.all():
            f.name(engines_in(li.inner_text()), "Pending engine choices")
    else:
        f.add("tied", 0, "no Pending engine choices section")

    removed: set[str] = set()
    summary = page.locator("p.exec")
    if summary.count():
        removed |= _other_targets(norm(summary.inner_text()))
    note = page.locator("[data-figure=eliminated-engines]")
    if note.count():
        removed |= engines_in(re.sub(r"see Engines not used below.*", "", note.inner_text()))
        if "see Engines not used below" in note.inner_text():
            f.add(
                "engines_not_used_section",
                page.locator("[data-figure=engines-not-used]").count() > 0,
                "Already removed note points to Engines not used",
            )
    not_used = page.locator("[data-figure=engines-not-used] li[data-engine]")
    removed |= {li.get_attribute("data-engine") for li in not_used.all()}
    if removed or summary.count():
        f.add("removed", frozenset(removed), "summary, Already removed note and Engines not used")
    return f


# --------------------------------------------------------------------------
# Interactive analysis report (the UI export, read in a browser)
# --------------------------------------------------------------------------

_ANALYSIS_WAVE = re.compile(
    r"^([\d,]+) queries · [\d.]+% of (?:calls|query patterns)"
    r"(?: · ([\d,]+) served, ([\d,]+) move| · ([\d,]+) source tables"
    r"(?: and views \(([\d,]+) tables?, ([\d,]+) views?\))?)?$"
)


def analysis_figures(page: Any) -> Figures:
    """Figures the interactive analysis report shows, read from the rendered page."""
    f = Figures("analysis report")

    def stat(label: str):
        return page.locator(
            ".stat-card", has=page.locator(f".stat-label:text-is('{label}')")
        ).locator(".stat-value")

    m = re.match(r"^\$([\d,.]+)/mo$", norm(stat("Projected Cost").inner_text()))
    f.add("cost:total", _usd(m.group(1)) if m else UNREADABLE, "Projected Cost stat")
    badges = stat("Target Engines").locator("[data-engine]").all()
    for b in badges:
        f.name(b.get_attribute("data-engine"), "Target Engines")
    f.add("engine_count", len(badges), "Target Engines stat")
    cache_badges = [
        b.get_attribute("data-engine") for b in badges if "(cache layer)" in b.inner_text()
    ]
    cache = stat("Cache Layer")
    if cache.count() and cache_badges:
        m = re.search(r"([\d,]+) cached reads?", cache.inner_text())
        if m:
            f.add(f"cached_reads:{cache_badges[0]}", _int(m.group(1)), "Cache Layer stat")

    for card in page.locator("#migration-roadmap-container [data-wave]").all():
        n = card.get_attribute("data-wave")
        where = f"roadmap wave {n}"
        for b in card.locator("[data-engine]").all():
            f.name(b.get_attribute("data-engine"), where)
        stats = norm(card.locator("[data-figure=wave-stats]").inner_text())
        m = _ANALYSIS_WAVE.match(stats)
        if not m:
            f.add(f"wave{n}:unparsed", stats, where)
            continue
        f.add(f"wave{n}:queries", _int(m.group(1)), where)
        if m.group(2):
            served, moving = _int(m.group(2)), _int(m.group(3))
            f.add(f"wave{n}:served", served, where)
            f.add(f"wave{n}:moving", moving, where)
            f.add(f"wave{n}:synced", served - moving, where)
        elif m.group(4):
            f.add(f"wave{n}:served", _int(m.group(4)), where)
            if m.group(5):
                f.add(f"wave{n}:tables_only", _int(m.group(5)), where)
                f.add(f"wave{n}:views", _int(m.group(6)), where)

    for card in page.locator("#cost-breakdown-container .stat-card").all():
        eng = card.locator("[data-engine]").first.get_attribute("data-engine")
        f.name(eng, "Cost Breakdown")
        m = re.search(r"\$([\d,]+\.\d\d)", card.inner_text())
        if m:
            f.add(f"cost:{eng}", _usd(m.group(1)), "Cost Breakdown card")

    risk_text = page.locator("#risk-assessment-container").inner_text()
    total = 0
    for sev, k in re.findall(r"^(CRITICAL|HIGH|MEDIUM|LOW|UNRATED) \((\d+)\)$", risk_text, re.M):
        f.add(f"risks:{sev}", int(k), "Risk Assessment severity heading")
        total += int(k)
    f.add("risks:total", total, "Risk Assessment severity headings")
    for b in page.locator("#risk-assessment-container > div [data-engine]").all():
        f.name(b.get_attribute("data-engine"), "Risk Assessment")

    for label in page.locator("#query-flow-container [data-figure=engine-queries]").all():
        eng = label.get_attribute("data-engine")
        f.name(eng, "Query Flow")
        m = re.search(r"\(([\d,]+) quer(?:y|ies), [\d.]+%\)$", norm(label.text_content() or ""))
        f.add(f"queries:{eng}", _int(m.group(1)) if m else UNREADABLE, "Query Flow label")

    summary = page.locator(".section-body").first
    f.add("removed", frozenset(_other_targets(norm(summary.inner_text()))), "summary")
    return f


# --------------------------------------------------------------------------
# Engineering report (Markdown)
# --------------------------------------------------------------------------


def md_plain(text: str) -> str:
    """Markdown source as the reader sees it: no backslash escapes or entities."""
    return html_lib.unescape(re.sub(r"\\([\\`*_{}\[\]()#+\-.!|<>~])", r"\1", text))


def _md_sections(md: str, level: int) -> list[tuple[str, str]]:
    """``(heading, body)`` for each heading of exactly ``level`` hashes."""
    marks = list(re.finditer(rf"^{'#' * level} (.+)$", md, re.M))
    out = []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(md)
        nxt = re.search(rf"^#{{1,{level}}} ", md[m.end() : end], re.M)
        body_end = m.end() + nxt.start() if nxt else end
        out.append((m.group(1).strip(), md[m.end() : body_end]))
    return out


def _md_table_rows(body: str) -> list[list[str]]:
    rows = []
    for line in body.splitlines():
        if line.startswith("|") and not re.match(r"^\|[-| ]+\|$", line):
            rows.append([c.strip() for c in re.split(r"(?<!\\)\|", line.strip())[1:-1]])
    return rows[1:]  # drop the header row


def engineering_figures(md_text: str) -> Figures:
    """Figures the engineering report shows."""
    f = Figures("engineering report")
    md = md_text
    plain = md_plain(md)
    h2 = dict(_md_sections(md, 2))

    targets: set[str] = set()
    for row in _md_table_rows(h2.get("Target engines", "")):
        eng = engine_id(md_plain(row[0]))
        targets.add(eng)
        f.name(eng, "Target engines table")
        m = re.match(r"^([\d,]+) cached reads?\b", md_plain(row[2]))
        if m:
            f.add(f"cached_reads:{eng}", _int(m.group(1)), "Target engines Workload")
    f.add("engine_count", len(targets), "Target engines table")

    m = re.search(r"^- Projected monthly \(target\): \$([\d,.]+)$", plain, re.M)
    f.add("cost:total", _usd(m.group(1)) if m else UNREADABLE, "Cost: Projected monthly")
    removed: set[str] = set()
    m = re.search(r"^- Already removed from this comparison: (.+)$", plain, re.M)
    if m:
        removed |= {engine_id(x) for x in re.findall(r"/mo \(([^)]+)\)", m.group(1))}

    for heading, body in _md_sections(md, 2):
        m = re.match(r"^Cache layer \((.+)\)$", heading)
        if m:
            k = re.search(r"fronts ([\d,]+) hot reads?", md_plain(body))
            if k:
                f.add(
                    f"cached_reads:{engine_id(m.group(1))}", _int(k.group(1)), "Cache layer section"
                )
        m = re.match(r"^Unresolved engine ties \(([\d,]+) quer(?:y|ies)\)$", heading)
        if m:
            f.add("tied", _int(m.group(1)), "Unresolved engine ties heading")
            rows = _md_table_rows(body)
            f.add("tied", len(rows), "Unresolved engine ties table")
            for row in rows:
                f.name(engines_in(md_plain(" ".join(row[1:]))), "Unresolved engine ties table")
        m = re.match(r"^Risk register \(([\d,]+)\)$", heading)
        if m:
            f.add("risks:total", _int(m.group(1)), "Risk register heading")
            sev = Counter(re.findall(r"^- \*\*RISK-\d+\*\* · (\w+) ·", body, re.M))
            for s in SEVERITIES:
                f.add(f"risks:{s}", sev.get(s, 0), "Risk register entries")
            for eng_heading, _ in _md_sections(body, 3):
                f.name(engine_id(md_plain(eng_heading)), "Risk register")
        if heading.startswith("Migration trade-offs"):
            for x in re.findall(r"Consolidated [\d,]+ quer(?:y|ies) from (\w+) →", md_plain(body)):
                if engine_id(x) not in targets:
                    removed.add(engine_id(x))
    if "Unresolved engine ties" not in " ".join(h2):
        f.add("tied", 0, "no Unresolved engine ties section")
    f.add("removed", frozenset(removed), "Already removed line and Migration trade-offs")

    queries_on_aurora = h2.get("Queries on Aurora by table", "")
    for m in re.finditer(r"^### (.+) \(([\d,]+) quer(?:y|ies)\)$", queries_on_aurora, re.M):
        f.add(f"queries:{engine_id(m.group(1))}", _int(m.group(2)), "Queries on Aurora by table")

    for heading, body in _md_sections(md, 3):
        m = re.match(r"^Wave (\d+): ", heading)
        if not m:
            continue
        n = m.group(1)
        where = f"roadmap wave {n}"
        text = md_plain(body)
        e = re.search(r"^- Engines: (.+)$", text, re.M)
        if e:
            f.name({engine_id(x) for x in e.group(1).split(",")}, where)
        c = re.search(
            r"^- all ([\d,]+) quer(?:y|ies) at cutover; ([\d,]+) \(", text, re.M
        ) or re.search(r"^- ()([\d,]+) quer(?:y|ies) \([\d.]+% of", text, re.M)
        if c:
            if c.group(1):
                f.add(f"wave{n}:cutover_queries", _int(c.group(1)), where)
            f.add(f"wave{n}:queries", _int(c.group(2)), where)
        st = re.search(
            r"^- ([\d,]+) source (?:tables and views(?: \(([\d,]+) tables?, ([\d,]+) views?\))?"
            r"|tables?): ",
            text,
            re.M,
        )
        if st:
            f.add(f"wave{n}:served", _int(st.group(1)), where)
            if st.group(2):
                f.add(f"wave{n}:tables_only", _int(st.group(2)), where)
                f.add(f"wave{n}:views", _int(st.group(3)), where)
        k = re.search(r"^- ([\d,]+) tables? kept in sync from their owner: ", text, re.M)
        if k:
            f.add(f"wave{n}:synced", _int(k.group(1)), where)
        g = re.search(r"^- [\d,]+ table groups?:\n((?:  - .+\n?)+)", text, re.M)
        if g:
            gq = sum(_int(x) for x in re.findall(r"^  - ([\d,]+) quer", g.group(1), re.M))
            f.add(f"wave{n}:group_queries", gq, where)
    return f


# --------------------------------------------------------------------------
# Executive deck (PPTX) and its PDF
# --------------------------------------------------------------------------


def deck_slide_texts(pptx_path: str) -> list[str]:
    """Each slide's text, shapes in deck order, table cells row by row, the
    same order the PDF draws them in."""
    from pptx import Presentation

    def walk(shapes: Any, out: list[str]) -> None:
        for sh in shapes:
            if getattr(sh, "shape_type", None) == 6:  # group
                walk(sh.shapes, out)
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    out.append(" ".join(c.text for c in row.cells))

    texts = []
    for slide in Presentation(pptx_path).slides:
        parts: list[str] = []
        walk(slide.shapes, parts)
        texts.append(norm(" ".join(parts)))
    return texts


def pdf_page_texts(pdf_path: str) -> list[str]:
    from pypdf import PdfReader

    return [norm(p.extract_text() or "") for p in PdfReader(pdf_path).pages]


_ROLES = "Retained|Migration target|Cache layer|Search read model|Evaluated|Assessed"


def deck_figures(slides: list[str], label: str) -> Figures:
    """Figures the executive deck shows, from its slide texts (the deck's own
    shapes, or the PDF drawn from them: one extractor reads both)."""
    f = Figures(label)
    text = " ".join(slides)

    m = re.search(r"([\d,]+) source tables? migrate\b", text)
    f.add("tables_migrate", _int(m.group(1)) if m else UNREADABLE, "Assessment Summary tile")
    m = re.search(r"overall risk, ([\d,]+) items?", text)
    if m:
        f.add("risks:total", _int(m.group(1)), "Assessment Summary risk tile")
    for m in re.finditer(
        rf"({_NAMES_RE}) (?:{_ROLES}) [\d.]+%(?: calls)? (\S.*?)(?= (?:{_NAMES_RE}) (?:{_ROLES}) |$)",
        text,
    ):
        eng = engine_id(m.group(1))
        f.name(eng, "engine tables")
        t = re.match(r"([\d,]+) source tables?\b", m.group(2))
        if t:
            f.add(f"tables:{eng}", _int(t.group(1)), "Assessment Summary engine table")
    for m in re.finditer(rf"({_NAMES_RE}) [\d.]+% · ([\d,]+) quer(?:y|ies) · \d+% fit", text):
        f.add(f"queries:{engine_id(m.group(1))}", _int(m.group(2)), "Workload Coverage")
        f.name(engine_id(m.group(1)), "Workload Coverage")
    for m in re.finditer(rf"({_NAMES_RE}) cache layer · ([\d,]+) cached reads?", text):
        f.add(f"cached_reads:{engine_id(m.group(1))}", _int(m.group(2)), "Workload Coverage")
        f.name(engine_id(m.group(1)), "Workload Coverage")

    m = re.search(r"Pick the engine for tied queries\? ([\d,]+) quer(?:y|ies)", text)
    f.add("tied", _int(m.group(1)) if m else 0, "Open Decisions tied-queries card")
    m = re.search(r"Scores tie between (.+?); no routing rule", text)
    if m:
        f.name(engines_in(m.group(1)), "Open Decisions tied-queries card")

    m = re.search(
        r"([\d,]+) risks?: ((?:[\d,]+ (?:CRITICAL|HIGH|MEDIUM|LOW|UNRATED)(?:, )?)+)", text
    )
    if m:
        f.add("risks:total", _int(m.group(1)), "Risk Profile subtitle")
        sev = {p.split()[1]: _int(p.split()[0]) for p in m.group(2).strip(", ").split(", ")}
        for s in SEVERITIES:
            f.add(f"risks:{s}", sev.get(s, 0), "Risk Profile subtitle")

    seq = next(
        (
            s
            for s in slides
            if s.startswith("Migration Sequencing") or " Migration Sequencing " in s
        ),
        "",
    )
    cards = re.split(r"\bWAVE (\d+) ", seq)
    for n, body in zip(cards[1::2], cards[2::2], strict=True):
        where = f"Migration Sequencing wave {n}"
        e = re.match(rf"((?:{_NAMES_RE})(?: \+ (?:{_NAMES_RE}))*) ", body)
        card_engines = [engine_id(x) for x in e.group(1).split(" + ")] if e else []
        f.name(set(card_engines), where)
        per_engine = re.findall(rf"({_NAMES_RE}) moves ([\d,]+) tables?", body)
        single = re.search(r"\bMoves ([\d,]+) tables?", body)
        if per_engine:
            for name, k in per_engine:
                f.add(f"wave{n}:moving:{engine_id(name)}", _int(k), where)
            f.add(f"wave{n}:moving", sum(_int(k) for _, k in per_engine), where)
        elif single:
            f.add(f"wave{n}:moving", _int(single.group(1)), where)
            if len(card_engines) == 1:
                f.add(f"wave{n}:moving:{card_engines[0]}", _int(single.group(1)), where)
        sync = re.search(r"Kept in sync from [^:]+: ([\d,]+) tables?", body)
        if sync:
            f.add(f"wave{n}:synced", _int(sync.group(1)), where)
        elif per_engine or single:
            f.add(f"wave{n}:synced", 0, where)

    summary = re.search(r"EXECUTIVE SUMMARY (.+?)(?= Engine Role Workload Scope|$)", text)
    if summary:
        f.add("removed", frozenset(_other_targets(summary.group(1))), "executive summary")
    return f
