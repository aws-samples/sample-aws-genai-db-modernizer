"""Unit tests for the orchestrator publish wiring (Increment 2).

``_publish_synthesis_deliverables`` (in ``tools``) is tested with a fake store
and a patched ``publish`` to confirm: three text deliverables plus the executive
deck and its PDF are written to S3 under the canonical artifact names, all five
deliverables are published as CUSTOMER_OUTPUT each carrying an explicit download
filename, and the whole step is non-fatal when the report cannot be read.

``_FakeStore`` implements ``write_bytes`` deliberately. While it did not, the
executive-summary block raised ``AttributeError`` straight into its own
``except`` and the fifth deliverable was never exercised at all.

Exercised against the same committed fixture used by the renderer tests (the real
report from the deployed ``v2-e2e-09`` run) — see
``tests/unit/report/test_report_renderers.py`` for the pure-renderer coverage.
"""

from __future__ import annotations

import json
import re
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import Mock, patch

import pytest

from src.atx_orchestrator.runtime import graph_transport
from src.atx_orchestrator.tools import run_synthesis_via_a2a
from src.report.deliverables import DeliverableSet

FIXTURE = Path(__file__).resolve().parents[1] / "report" / "fixtures" / "e2e09_report.json"


@pytest.fixture(scope="module")
def report() -> dict:
    data: dict = json.loads(FIXTURE.read_text())
    return data


class _FakeStore:
    def __init__(self, data: dict) -> None:
        self.data: dict[str, dict] = dict(data)
        self.text_writes: dict[str, tuple[str, str]] = {}
        self.byte_writes: dict[str, bytes] = {}

    def read_json(self, path: str) -> dict:
        return self.data[path]

    def exists(self, path: str) -> bool:
        return path in self.data

    def list_prefix(self, prefix: str) -> list[str]:
        return [k for k in self.data if k.startswith(prefix)]

    def write_text(self, path: str, content: str, content_type: str = "text/plain") -> None:
        self.text_writes[path] = (content, content_type)

    # The executive summary writes binary. Without this the whole block raises
    # AttributeError into its own ``except`` and the fifth deliverable silently
    # never renders -- which is exactly how it went untested before.
    def write_bytes(self, path: str, content: bytes, content_type: str = "") -> None:
        self.byte_writes[path] = content


class _JsonOnlyStore(_FakeStore):
    """Models the ATX artifact store: JSON-only, raises on text/bytes writes.

    Used to prove that when the durable "system of record" copy cannot be written
    (ATX backend), the synthesis publish path still reaches artifacts.publish with
    every deliverable — the customer must still get their downloadable reports.
    """

    def write_text(self, path: str, content: str, content_type: str = "text/plain") -> None:
        raise NotImplementedError("ATX artifact store is JSON-only")

    def write_bytes(self, path: str, content: bytes, content_type: str = "") -> None:
        raise NotImplementedError("ATX artifact store is JSON-only")


class TestSynthesisDeliverables:
    KEY = "discourse/job-x/synthesis/v1/report.json"

    def test_writes_five_files_and_publishes_five(self, report: dict) -> None:
        payload = {"response": {"report_artifact": self.KEY, "engines_ranked": 5}}
        store = _FakeStore({self.KEY: report})
        captured: dict = {}
        with (
            patch("src.atx_orchestrator.tools.invoke_and_wait", return_value=payload),
            patch("src.atx_orchestrator.tools._make_store", return_value=store),
            patch(
                "src.atx_orchestrator.runtime.artifacts.publish",
                side_effect=lambda items: captured.update({"items": items}) or {},
            ),
        ):
            out = run_synthesis_via_a2a("job-x", "discourse")

        # payload flows through unchanged
        assert json.loads(out)["response"]["report_artifact"] == self.KEY

        # provenance stems are date-stamped; derive the day rather than freezing time
        day = datetime.now(UTC).strftime("%Y%m%d")

        # exactly the three rendered deliverables written to S3 (report.json untouched)
        keys = list(store.text_writes)
        assert len(keys) == 3
        # canonical stem: {database}_{artifact}_{job8}_{YYYYMMDD}.{ext}
        assert any(re.search(r"/discourse_decision-report_job-x_\d{8}\.html$", k) for k in keys)
        assert any(re.search(r"/discourse_engineering-report_job-x_\d{8}\.md$", k) for k in keys)
        assert any(re.search(r"/discourse_analysis-report_job-x_\d{8}\.html$", k) for k in keys)
        assert self.KEY not in keys

        # the executive summary writes both halves as binary under fixed names --
        # the deck stays on S3 as the editable source, only the PDF is registered
        assert sorted(k.rsplit("/", 1)[-1] for k in store.byte_writes) == [
            "summary-executive-report.pdf",
            "summary-executive-report.pptx",
        ]
        assert (
            store.byte_writes[f"{self.KEY.rsplit('/', 1)[0]}/summary-executive-report.pdf"][:5]
            == b"%PDF-"
        )

        # five published, in order, all CUSTOMER_OUTPUT
        items = captured["items"]
        assert [it[1] for it in items] == ["HTML", "MARKDOWN", "JSON", "HTML", "PDF"]
        assert [it[2] for it in items] == [
            "Decision Report — discourse",
            "Engineering Report — discourse",
            "Assessment Data (raw) — discourse",
            "Interactive Analysis Report — discourse",
            "Executive Summary Report — discourse",
        ]
        assert {it[3] for it in items} == {"CUSTOMER_OUTPUT"}

        # Every item carries an explicit download filename (5th element). Without it
        # publish() cannot set fileMetadata.path and the customer's download is named
        # after the artifact UUID -- a regression no other assertion here would catch.
        assert all(len(it) == 5 for it in items)
        assert items[-1][4] == "summary-executive-report.pdf"
        assert [it[4] for it in items[:4]] == [
            f"discourse_decision-report_job-x_{day}.html",
            f"discourse_engineering-report_job-x_{day}.md",
            f"discourse_assessment-data_job-x_{day}.json",
            f"discourse_analysis-report_job-x_{day}.html",
        ]

        # Waves alignment (ADR-029): the executive deck reframes the raw
        # target-engine count as migration waves ("why move to 5 databases?" is
        # the most common objection). Parse the rendered .pptx and confirm the
        # summary tile dropped "target engines" for "migration waves".
        import io as _io

        from pptx import Presentation

        deck_key = next(k for k in store.byte_writes if k.endswith(".pptx"))
        prs = Presentation(_io.BytesIO(store.byte_writes[deck_key]))
        deck_text = " ".join(
            run.text
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_text_frame
            for para in shape.text_frame.paragraphs
            for run in para.runs
        )
        assert "migration waves" in deck_text
        assert "target engines" not in deck_text

    def test_publishes_all_five_on_json_only_store(self, report: dict) -> None:
        """On the JSON-only ATX backend the durable store copies raise
        NotImplementedError, but publish() must still deliver all five reports --
        publish is the actual customer-facing delivery, the store copy is only a
        system-of-record duplicate that backend has no place for."""
        payload = {"response": {"report_artifact": self.KEY, "engines_ranked": 5}}
        store = _JsonOnlyStore({self.KEY: report})
        captured: dict = {}
        with (
            patch("src.atx_orchestrator.tools.invoke_and_wait", return_value=payload),
            patch("src.atx_orchestrator.tools._make_store", return_value=store),
            patch(
                "src.atx_orchestrator.runtime.artifacts.publish",
                side_effect=lambda items: captured.update({"items": items}) or {},
            ),
        ):
            run_synthesis_via_a2a("job-x", "discourse")

        # No durable copies landed (the store rejected them) ...
        assert store.text_writes == {}
        assert store.byte_writes == {}
        # ... but all five deliverables were still published to the customer.
        items = captured["items"]
        assert [it[1] for it in items] == ["HTML", "MARKDOWN", "JSON", "HTML", "PDF"]
        assert {it[3] for it in items} == {"CUSTOMER_OUTPUT"}

    def test_non_fatal_when_report_unreadable(self) -> None:
        payload = {"response": {"report_artifact": "missing/key.json"}}
        store = _FakeStore({})  # read_json raises KeyError
        with (
            patch("src.atx_orchestrator.tools.invoke_and_wait", return_value=payload),
            patch("src.atx_orchestrator.tools._make_store", return_value=store),
        ):
            out = run_synthesis_via_a2a("job-x", "discourse")
        assert json.loads(out)["response"]["report_artifact"] == "missing/key.json"

    def test_non_fatal_when_no_report_artifact(self) -> None:
        payload = {"response": {"engines_ranked": 5}}
        with (
            patch("src.atx_orchestrator.tools.invoke_and_wait", return_value=payload),
            patch("src.atx_orchestrator.tools._make_store"),
        ):
            out = run_synthesis_via_a2a("job-x", "discourse")
        # No report_artifact -> deliverables publishing is skipped; the payload
        # flows through unchanged. (A store is constructed once for version
        # resolution, which is best-effort and never fatal.)
        assert json.loads(out) == payload

    def test_pins_the_graph_fetcher_and_assignment_version(self, report: dict) -> None:
        """render_deliverables must be called with the ATX graph transport's
        download_graph -- not the local-rebuild default -- and the
        assignment_version resolved from the payload. The ATX store is
        JSON-only; the graph travels as a published artifact, so silently
        falling back to the local default here would rebuild from JSON
        contracts instead of reading what synthesis actually published."""
        payload = {"response": {"report_artifact": self.KEY, "engines_ranked": 5}}
        store = _FakeStore({self.KEY: report})
        mock_render = Mock(return_value=DeliverableSet())
        with (
            patch("src.atx_orchestrator.tools.invoke_and_wait", return_value=payload),
            patch("src.atx_orchestrator.tools._make_store", return_value=store),
            patch("src.report.deliverables.render_deliverables", mock_render),
            patch("src.atx_orchestrator.runtime.artifacts.publish", return_value={}),
        ):
            run_synthesis_via_a2a("job-x", "discourse")

        assert mock_render.called
        _, kwargs = mock_render.call_args
        assert kwargs["graph_fetcher"] is graph_transport.download_graph
        assert kwargs["assignment_version"] == 1
