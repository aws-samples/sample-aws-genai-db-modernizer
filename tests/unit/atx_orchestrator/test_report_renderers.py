"""Unit tests for the customer-facing report renderers and the orchestrator
publish wiring (Increment 2).

The renderers (``render_decision_report_html``, ``render_engineering_report_md``)
and their helpers are pure functions of the synthesis ``report.json``. They are
exercised against a committed fixture — the real report from the deployed
``v2-e2e-09`` run — so the assertions check genuine reconciliation
($814.12 / 100% workload / 92 migrated tables) and the executive/engineering
content split, not a hand-built stub.

``_publish_synthesis_deliverables`` (in ``tools``) is tested with a fake store
and a patched ``publish`` to confirm: three text deliverables plus the executive
deck and its PDF are written to S3 under the canonical artifact names, all five
deliverables are published as CUSTOMER_OUTPUT each carrying an explicit download
filename, and the whole step is non-fatal when the report cannot be read.

``_FakeStore`` implements ``write_bytes`` deliberately. While it did not, the
executive-summary block raised ``AttributeError`` straight into its own
``except`` and the fifth deliverable was never exercised at all.
"""

from __future__ import annotations

import json
import re
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import patch

import pytest

from src.atx_orchestrator.runtime import artifacts
from src.atx_orchestrator.tools import run_synthesis_via_a2a

FIXTURE = Path(__file__).parent / "fixtures" / "e2e09_report.json"


@pytest.fixture(scope="module")
def report() -> dict:
    data: dict = json.loads(FIXTURE.read_text())
    return data


# =============================================================================
# Decision Report (HTML, executive)


class TestDecisionReport:
    def test_renders_html_document(self, report: dict) -> None:
        h = artifacts.render_decision_report_html(report)
        low = h.lstrip().lower()
        assert low.startswith("<!doctype html") or low.startswith("<html")
        assert len(h) > 3000

    def test_offline_no_external_loads(self, report: dict) -> None:
        """Self-contained: no CDN/CSS/JS/font loads. The only URI allowed is the
        inline SVG xmlns."""
        h = artifacts.render_decision_report_html(report)
        assert "https://" not in h
        assert set(re.findall(r"http://[^\s\"'>]+", h)) <= {"http://www.w3.org/2000/svg"}
        assert "<script" not in h.lower()
        assert "<link" not in h.lower()
        assert "cdn" not in h.lower()

    def test_five_engine_reconciliation(self, report: dict) -> None:
        h = artifacts.render_decision_report_html(report)
        for eng in ("aurora", "elasticache", "documentdb", "dynamodb", "opensearch"):
            assert eng in h.lower()
        assert "814.12" in h
        assert "92" in h

    def test_architecture_engines_all_five_with_roles(self, report: dict) -> None:
        engines = artifacts._architecture_engines(report)
        assert len(engines) == 5
        total_workload = sum(e.get("workload", 0) for e in engines)
        assert 99.0 <= total_workload <= 101.0
        total_cost = sum(e.get("cost", 0) for e in engines)
        assert abs(total_cost - 814.12) < 1.0
        roles = {e["role"] for e in engines}
        # at least the three role kinds present: retained, migration target, cache
        assert any("etain" in r for r in roles)
        assert any("igration" in r for r in roles)
        assert any("ache" in r for r in roles)

    def test_no_per_risk_list_no_tradeoffs(self, report: dict) -> None:
        h = artifacts.render_decision_report_html(report)
        assert "Risk posture" in h
        assert "Key trade-offs" not in h
        # no per-risk cards in the body
        assert 'class="risk ' not in h

    def test_risk_posture_counts_and_strategies(self, report: dict) -> None:
        h = artifacts.render_decision_report_html(report)
        assert "22" in h
        assert "Mitigation strategies" in h

    def test_no_empty_unknown_risks(self, report: dict) -> None:
        assert "unknown:" not in artifacts.render_decision_report_html(report)


# =============================================================================
# Engineering Report (Markdown, build team)


class TestEngineeringReport:
    def test_renders_markdown(self, report: dict) -> None:
        m = artifacts.render_engineering_report_md(report)
        assert m.startswith("# Database Modernization")
        assert len(m) > 5000

    def test_risk_register_present_and_filtered(self, report: dict) -> None:
        m = artifacts.render_engineering_report_md(report)
        assert "## Risk register (22)" in m
        # the six malformed empty risks ([engine] unknown: with no body) are dropped.
        # (Note "unknown:" can still appear as a sub-type label on a KEPT risk that
        # has a real body after it, so we check the empty IDs, not the literal.)
        for rid in ("RISK-007", "RISK-008", "RISK-012", "RISK-013", "RISK-014", "RISK-015"):
            assert rid not in m

    def test_tradeoffs_by_engine(self, report: dict) -> None:
        m = artifacts.render_engineering_report_md(report)
        assert "## Migration trade-offs (44)" in m
        assert "### documentdb" in m

    def test_has_mermaid_fences(self, report: dict) -> None:
        m = artifacts.render_engineering_report_md(report)
        assert "```mermaid" in m


# =============================================================================
# Empty-risk filter helpers


class TestRiskFilter:
    def test_filters_six_empty_risks(self, report: dict) -> None:
        risks = report["risk_assessment"]["risks"]
        kept = [r for r in risks if artifacts._risk_has_content(r.get("description"))]
        assert len(risks) == 28
        assert len(kept) == 22

    def test_engine_and_body_parse(self) -> None:
        eng, body = artifacts._risk_engine_and_body("[documentdb] Queries joining 3+ tables")
        assert eng == "documentdb"
        assert body == "Queries joining 3+ tables"

    def test_has_content_predicate(self) -> None:
        assert artifacts._risk_has_content("[elasticache] unknown: ") is False
        assert artifacts._risk_has_content("[documentdb] a real risk") is True
        assert artifacts._risk_has_content("") is False


# =============================================================================
# Engine roles — an engine that won no queries is not "Retained"


@pytest.fixture
def zero_workload_report(report: dict) -> dict:
    """The real report, reshaped so DocumentDB is the engine that won nothing.

    Reproduces the observed defect: triage selected DocumentDB, analysis scored it 56%,
    then the assignment routed every query elsewhere, so schema design was skipped.
    """
    rep: dict = json.loads(json.dumps(report))
    pct = {
        "aurora_postgresql": 49.3,
        "elasticache": 29.4,
        "dynamodb": 15.1,
        "opensearch": 6.1,
        "documentdb": 0.0,
    }
    for r in rep["ranking"]:
        r["workload_percent"] = pct.get(r["target"])
    rep["recommended_architecture"]["databases"] = [
        d for d in rep["recommended_architecture"]["databases"] if d["service"] != "documentdb"
    ]
    rep.setdefault("schema_designs", {})["documentdb"] = {"status": "skipped"}
    return rep


class TestEngineRole:
    def test_zero_workload_is_evaluated_not_retained(self) -> None:
        """The defect: a skipped design used to imply "Retained" regardless of workload."""
        assert (
            artifacts._engine_role("documentdb", set(), {"documentdb": {"status": "skipped"}}, 0.0)
            == "Evaluated"
        )
        assert artifacts._engine_role("documentdb", set(), {}, None) == "Evaluated"

    def test_workload_carrying_engine_with_no_design_is_still_retained(self) -> None:
        """Guards against over-correcting: the source relational core keeps Retained."""
        role = artifacts._engine_role(
            "aurora_postgresql", set(), {"aurora_postgresql": {"status": "not_available"}}, 49.3
        )
        assert role == "Retained"

    def test_cache_and_migration_target_keep_their_role_at_zero_workload(self) -> None:
        """Branch-order guard: the workload test must not outrank these two."""
        assert artifacts._engine_role("elasticache", set(), {}, 0.0) == "Cache layer"
        assert artifacts._engine_role("dynamodb", {"dynamodb"}, {}, 0.0) == "Migration target"

    def test_scope_reads_no_queries_assigned(self, zero_workload_report: dict) -> None:
        rows = {e["engine"]: e for e in artifacts._architecture_engines(zero_workload_report)}
        assert rows["documentdb"]["role"] == "Evaluated"
        assert rows["documentdb"]["scope"] == "no queries assigned"
        # the engines that do the work are unaffected
        assert rows["aurora_postgresql"]["role"] == "Retained"
        assert rows["elasticache"]["role"] == "Cache layer"

    def test_evaluated_engine_excluded_from_retained_narrative(
        self, zero_workload_report: dict
    ) -> None:
        """The HTML claimed DocumentDB was "retained as the relational core"."""
        engines = artifacts._architecture_engines(zero_workload_report)
        assert [e["engine"] for e in engines if e["role"] == "Retained"] == ["aurora_postgresql"]

    def test_evaluated_engine_excluded_from_wave_one(self, zero_workload_report: dict) -> None:
        """Wave 1 is built from ``no_move``, and its confidence is a ``min()`` over members.

        An engine the assignment routed nothing to must not be named in the wave nor be
        eligible to set its confidence floor. On the observed ``discourse`` report the
        printed figure does not move (ElastiCache independently sits at the same 56%), so
        this asserts membership rather than the number — the number is only distorted when
        the Evaluated engine happens to be the sole minimum, and asserting it here would
        encode a coincidence of that one report.
        """
        engines = artifacts._architecture_engines(zero_workload_report)
        no_move = [e for e in engines if e["role"] in ("Retained", "Cache layer")]
        assert [e["engine"] for e in no_move] == ["aurora_postgresql", "elasticache"]
        assert "documentdb" not in {e["engine"] for e in no_move}

    def test_footer_names_and_percentage_agree(self, zero_workload_report: dict) -> None:
        """Tested positively on role, so an Evaluated engine is not named as "keeping" workload."""
        engines = artifacts._architecture_engines(zero_workload_report)
        kept = [e for e in engines if e["role"] in ("Retained", "Cache layer")]
        assert [e["engine"] for e in kept] == ["aurora_postgresql", "elasticache"]
        assert abs(sum(e["workload"] for e in kept) - 78.7) < 0.05

    def test_architecture_svg_renders_the_evaluated_engine_muted(
        self, zero_workload_report: dict
    ) -> None:
        svg = artifacts.architecture_svg(zero_workload_report)
        assert artifacts._ROLE_STROKE["Evaluated"] in svg
        assert ">None<" not in svg


# =============================================================================
# Orchestrator publish wiring


class _FakeStore:
    def __init__(self, data: dict) -> None:
        self.data: dict[str, dict] = dict(data)
        self.text_writes: dict[str, tuple[str, str]] = {}
        self.byte_writes: dict[str, bytes] = {}

    def read_json(self, path: str) -> dict:
        return self.data[path]

    def exists(self, path: str) -> bool:
        return path in self.data

    def list_prefix(self, prefix: str) -> list[str]:
        return [k for k in self.data if k.startswith(prefix)]

    def write_text(self, path: str, content: str, content_type: str = "text/plain") -> None:
        self.text_writes[path] = (content, content_type)

    # The executive summary writes binary. Without this the whole block raises
    # AttributeError into its own ``except`` and the fifth deliverable silently
    # never renders -- which is exactly how it went untested before.
    def write_bytes(self, path: str, content: bytes, content_type: str = "") -> None:
        self.byte_writes[path] = content


class _JsonOnlyStore(_FakeStore):
    """Models the ATX artifact store: JSON-only, raises on text/bytes writes.

    Used to prove that when the durable "system of record" copy cannot be written
    (ATX backend), the synthesis publish path still reaches artifacts.publish with
    every deliverable — the customer must still get their downloadable reports.
    """

    def write_text(self, path: str, content: str, content_type: str = "text/plain") -> None:
        raise NotImplementedError("ATX artifact store is JSON-only")

    def write_bytes(self, path: str, content: bytes, content_type: str = "") -> None:
        raise NotImplementedError("ATX artifact store is JSON-only")


class TestSynthesisDeliverables:
    KEY = "discourse/job-x/synthesis/v1/report.json"

    def test_writes_five_files_and_publishes_five(self, report: dict) -> None:
        payload = {"response": {"report_artifact": self.KEY, "engines_ranked": 5}}
        store = _FakeStore({self.KEY: report})
        captured: dict = {}
        with (
            patch("src.atx_orchestrator.tools.invoke_and_wait", return_value=payload),
            patch("src.atx_orchestrator.tools._make_store", return_value=store),
            patch(
                "src.atx_orchestrator.runtime.artifacts.publish",
                side_effect=lambda items: captured.update({"items": items}) or {},
            ),
        ):
            out = run_synthesis_via_a2a("job-x", "discourse")

        # payload flows through unchanged
        assert json.loads(out)["response"]["report_artifact"] == self.KEY

        # provenance stems are date-stamped; derive the day rather than freezing time
        day = datetime.now(UTC).strftime("%Y%m%d")

        # exactly the three rendered deliverables written to S3 (report.json untouched)
        keys = list(store.text_writes)
        assert len(keys) == 3
        # canonical stem: {database}_{artifact}_{job8}_{YYYYMMDD}.{ext}
        assert any(re.search(r"/discourse_decision-report_job-x_\d{8}\.html$", k) for k in keys)
        assert any(re.search(r"/discourse_engineering-report_job-x_\d{8}\.md$", k) for k in keys)
        assert any(re.search(r"/discourse_analysis-report_job-x_\d{8}\.html$", k) for k in keys)
        assert self.KEY not in keys

        # the executive summary writes both halves as binary under fixed names --
        # the deck stays on S3 as the editable source, only the PDF is registered
        assert sorted(k.rsplit("/", 1)[-1] for k in store.byte_writes) == [
            "summary-executive-report.pdf",
            "summary-executive-report.pptx",
        ]
        assert (
            store.byte_writes[f"{self.KEY.rsplit('/', 1)[0]}/summary-executive-report.pdf"][:5]
            == b"%PDF-"
        )

        # five published, in order, all CUSTOMER_OUTPUT
        items = captured["items"]
        assert [it[1] for it in items] == ["HTML", "MARKDOWN", "JSON", "HTML", "PDF"]
        assert [it[2] for it in items] == [
            "Decision Report — discourse",
            "Engineering Report — discourse",
            "Assessment Data (raw) — discourse",
            "Interactive Analysis Report — discourse",
            "Executive Summary Report — discourse",
        ]
        assert {it[3] for it in items} == {"CUSTOMER_OUTPUT"}

        # Every item carries an explicit download filename (5th element). Without it
        # publish() cannot set fileMetadata.path and the customer's download is named
        # after the artifact UUID -- a regression no other assertion here would catch.
        assert all(len(it) == 5 for it in items)
        assert items[-1][4] == "summary-executive-report.pdf"
        assert [it[4] for it in items[:4]] == [
            f"discourse_decision-report_job-x_{day}.html",
            f"discourse_engineering-report_job-x_{day}.md",
            f"discourse_assessment-data_job-x_{day}.json",
            f"discourse_analysis-report_job-x_{day}.html",
        ]

        # Waves alignment (ADR-029): the executive deck reframes the raw
        # target-engine count as migration waves ("why move to 5 databases?" is
        # the most common objection). Parse the rendered .pptx and confirm the
        # summary tile dropped "target engines" for "migration waves".
        import io as _io

        from pptx import Presentation

        deck_key = next(k for k in store.byte_writes if k.endswith(".pptx"))
        prs = Presentation(_io.BytesIO(store.byte_writes[deck_key]))
        deck_text = " ".join(
            run.text
            for slide in prs.slides
            for shape in slide.shapes
            if shape.has_text_frame
            for para in shape.text_frame.paragraphs
            for run in para.runs
        )
        assert "migration waves" in deck_text
        assert "target engines" not in deck_text

    def test_publishes_all_five_on_json_only_store(self, report: dict) -> None:
        """On the JSON-only ATX backend the durable store copies raise
        NotImplementedError, but publish() must still deliver all five reports —
        publish is the actual customer-facing delivery, the store copy is only a
        system-of-record duplicate that backend has no place for."""
        payload = {"response": {"report_artifact": self.KEY, "engines_ranked": 5}}
        store = _JsonOnlyStore({self.KEY: report})
        captured: dict = {}
        with (
            patch("src.atx_orchestrator.tools.invoke_and_wait", return_value=payload),
            patch("src.atx_orchestrator.tools._make_store", return_value=store),
            patch(
                "src.atx_orchestrator.runtime.artifacts.publish",
                side_effect=lambda items: captured.update({"items": items}) or {},
            ),
        ):
            run_synthesis_via_a2a("job-x", "discourse")

        # No durable copies landed (the store rejected them) ...
        assert store.text_writes == {}
        assert store.byte_writes == {}
        # ... but all five deliverables were still published to the customer.
        items = captured["items"]
        assert [it[1] for it in items] == ["HTML", "MARKDOWN", "JSON", "HTML", "PDF"]
        assert {it[3] for it in items} == {"CUSTOMER_OUTPUT"}

    def test_non_fatal_when_report_unreadable(self) -> None:
        payload = {"response": {"report_artifact": "missing/key.json"}}
        store = _FakeStore({})  # read_json raises KeyError
        with (
            patch("src.atx_orchestrator.tools.invoke_and_wait", return_value=payload),
            patch("src.atx_orchestrator.tools._make_store", return_value=store),
        ):
            out = run_synthesis_via_a2a("job-x", "discourse")
        assert json.loads(out)["response"]["report_artifact"] == "missing/key.json"

    def test_non_fatal_when_no_report_artifact(self) -> None:
        payload = {"response": {"engines_ranked": 5}}
        with (
            patch("src.atx_orchestrator.tools.invoke_and_wait", return_value=payload),
            patch("src.atx_orchestrator.tools._make_store"),
        ):
            out = run_synthesis_via_a2a("job-x", "discourse")
        # No report_artifact -> deliverables publishing is skipped; the payload
        # flows through unchanged. (A store is constructed once for version
        # resolution, which is best-effort and never fatal.)
        assert json.loads(out) == payload
