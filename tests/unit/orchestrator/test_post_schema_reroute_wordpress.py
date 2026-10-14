"""#713 review round 2 repro: a PE flag on wordpress moves one query after approval.

The deterministic wordpress pipeline runs to the approved assignment (v2), the
customer approves it, DynamoDB gets a design (a stand-in designer that designs
exactly the engine's assigned queries), and DynamoDB's PE review flags
``d72966f8...``. The post-schema router moves it to Aurora MySQL (54 against 69)
and writes v3. Then synthesis and the reports run on v3, and every deliverable
must agree with the v3 assignment: the moved query appears once, on Aurora
MySQL; DynamoDB's counts match its own v3 queries; the move is listed as made
after approval, with the override path; and v3 never claims the customer
approved it.
"""

from __future__ import annotations

import copy
import json
import re
from unittest.mock import patch

import pytest

from src.agents.referee.assignment_overrides import mark_assignment_customer_approved
from src.orchestrator.local_orchestrator import LocalOrchestrator
from src.storage.local_store import LocalArtifactStore
from tests.e2e.pipeline import REPO, PipelineResult, _run, run_pipeline

DB, JOB = "wordpress", "job-reroute"
FLAGGED = "d72966f8"
TEMPLATE = json.loads(
    (
        REPO / "docs/examples/wordpress/load-test-results/schema-dynamodb/v2/schema_output.json"
    ).read_text()
)


def _fake_design(job_id, database_name, engine, store, assignment_version=0, llm_mode=""):
    """Designs exactly the engine's in-scope queries on that version (DynamoDB
    only; Aurora MySQL is left undesigned, so synthesis builds its groups from
    the assignment)."""
    if engine != "dynamodb":
        return
    assignment = store.read_json(
        f"{database_name}/{job_id}/assignment/v{assignment_version}/assignment.json"
    )
    queries = [
        qa["query_id"]
        for qa in assignment["query_assignments"]
        if qa["assigned_engine"] == engine and qa.get("in_scope", True)
    ]
    design = copy.deepcopy(TEMPLATE)
    design["job_id"] = job_id
    template = TEMPLATE["access_patterns"][0]
    design["access_patterns"] = [
        {**template, "pattern_id": f"DDB-AP-{i + 1}", "query_ids": [qid]}
        for i, qid in enumerate(queries)
    ]
    base = f"{database_name}/{job_id}/schema-{engine}/v{assignment_version}"
    store.write_json(f"{base}/schema_output.json", design)
    store.write_json(f"{base}/design_trace.json", {"trace": []})


@pytest.fixture(scope="module")
def rerouted(tmp_path_factory) -> dict:
    root = tmp_path_factory.mktemp("reroute") / "artifacts"
    run_pipeline("wordpress", root, JOB)
    store = LocalArtifactStore(base_dir=str(root))
    mark_assignment_customer_approved(store, DB, JOB)
    _fake_design(JOB, DB, "dynamodb", store, 2)
    v2 = store.read_json(f"{DB}/{JOB}/assignment/v2/assignment.json")
    flagged = next(
        qa["query_id"] for qa in v2["query_assignments"] if qa["query_id"].startswith(FLAGGED)
    )
    store.write_json(
        f"{DB}/{JOB}/schema-dynamodb/v2/design_trace.json",
        {
            "trace": [],
            "routing_flags": [
                {
                    "query_ids": [flagged],
                    "flagged_engine": "dynamodb",
                    "reason_category": "client_side_join",
                    "explanation": "needs a server-side join",
                }
            ],
        },
    )
    with patch("src.agents.schema_design.handler.run_schema_design_auto", side_effect=_fake_design):
        LocalOrchestrator(store, llm_mode="bedrock")._run_post_schema_routing(JOB, DB)
    result = PipelineResult(db=DB, job_id=JOB, artifact_root=root)
    common = ["--job-id", JOB, "--db", DB, "--artifact-root", str(root)]
    _run("synthesis", ["scripts/run_synthesis.py", *common, "--llm-mode", "none"], result)
    _run("report", ["scripts/run_report.py", *common], result)
    synthesis = root / DB / JOB / "synthesis" / "v3"
    return {
        "flagged": flagged,
        "v2": v2,
        "v3": store.read_json(f"{DB}/{JOB}/assignment/v3/assignment.json"),
        "report": json.loads((synthesis / "report.json").read_text()),
        "decision": next(synthesis.glob("*decision-report*.html")).read_text(),
        "engineering": next(synthesis.glob("*engineering-report*.md")).read_text(),
        "analysis": next(synthesis.glob("*analysis-report*.html")).read_text(),
        "deck": synthesis / "summary-executive-report.pptx",
    }


def _in_scope(assignment: dict, engine: str) -> int:
    return sum(
        1
        for qa in assignment["query_assignments"]
        if qa["assigned_engine"] == engine and qa.get("in_scope", True)
    )


class TestTheRerouteOnWordpress:
    def test_the_query_moves_to_the_best_kept_engine_on_a_new_version(self, rerouted):
        engine_of = {
            qa["query_id"]: qa["assigned_engine"] for qa in rerouted["v3"]["query_assignments"]
        }
        assert engine_of[rerouted["flagged"]] == "aurora_mysql"
        (reroute,) = rerouted["v3"]["post_schema_reroutes"]
        assert (reroute["score_before"], reroute["score_after"]) == (69, 54)

    def test_the_rerouted_version_never_claims_customer_approval(self, rerouted):
        assert rerouted["v2"]["status"] == "customer_approved"
        assert rerouted["v3"]["status"] == "rerouted_after_approval"
        assert rerouted["report"]["assignment_summary"]["status"] == "rerouted_after_approval"


class TestEveryDeliverableAgreesWithTheNewVersion:
    def test_the_moved_query_is_in_exactly_one_query_group_on_aurora(self, rerouted):
        groups = rerouted["report"]["query_groups"]
        holders = [
            g["engines"]
            for g in groups
            for ap in g["access_patterns"]
            if rerouted["flagged"] in (ap.get("query_ids") or [])
        ]
        assert holders == [["aurora_mysql"]]
        total = sum(len(ap.get("query_ids") or []) for g in groups for ap in g["access_patterns"])
        assert total == sum(
            1 for qa in rerouted["v3"]["query_assignments"] if qa.get("in_scope", True)
        )

    def test_dynamodb_counts_match_its_own_queries(self, rerouted):
        expected = _in_scope(rerouted["v3"], "dynamodb")
        assert expected == _in_scope(rerouted["v2"], "dynamodb") - 1
        text = " ".join(re.sub(r"<[^>]+>", " ", rerouted["decision"]).split())
        assert f"DynamoDB handles {expected} queries" in text
        assert f"{expected} in-scope access patterns)" in text
        assert f"DynamoDB: {_in_scope(rerouted['v2'], 'dynamodb')}" not in text

    def test_the_reroute_records_pass_the_customer_text_guard(self, rerouted):
        from src.tools.validation.customer_text_guard import scan_json

        findings = scan_json({"post_schema_reroutes": rerouted["report"]["post_schema_reroutes"]})
        assert findings == []

    def test_no_risk_or_unsupported_pattern_names_the_moved_query_on_dynamodb(self, rerouted):
        risks = rerouted["report"].get("risk_assessment", {}).get("risks", [])
        assert not [r for r in risks if rerouted["flagged"][:12] in json.dumps(r)]
        assert "Unsupported patterns" not in rerouted["engineering"].split("## Moved after")[0]

    def test_the_move_after_approval_and_the_override_path_are_in_every_report(self, rerouted):
        from pptx import Presentation

        from src.report.engine_choices import REROUTE_OVERRIDE_TEXT

        assert "Moved after your approval (1 query)" in rerouted["decision"]
        assert "## Moved after approval (1 query)" in rerouted["engineering"]
        assert REROUTE_OVERRIDE_TEXT in rerouted["decision"]
        assert REROUTE_OVERRIDE_TEXT in rerouted["engineering"]
        deck_text = " ".join(
            run.text
            for slide in Presentation(str(rerouted["deck"])).slides
            for shape in slide.shapes
            if shape.has_text_frame
            for paragraph in shape.text_frame.paragraphs
            for run in paragraph.runs
        )
        assert "Keep the queries moved after approval?" in deck_text
        # The interactive analysis report carries the list and renders it.
        assert '"post_schema_reroutes"' in rerouted["analysis"]
        assert rerouted["flagged"] in rerouted["analysis"]
        assert "buildPostSchemaReroutes();" in rerouted["analysis"]
        assert rerouted["report"]["post_schema_reroutes"][0]["approved_version"] == 2


@pytest.fixture(scope="module")
def overridden(tmp_path_factory) -> dict:
    """#713 review round 3 probe: v3 moves the query to Aurora MySQL; the
    customer moves it back to DynamoDB at the review (v4, customer_override);
    DynamoDB is redesigned on v4 and its PE review flags the query again; the
    router runs again. The customer's choice must hold: no v5."""
    from src.agents.referee.assignment_overrides import (
        QueryOverrideInput,
        apply_assignment_overrides,
    )

    root = tmp_path_factory.mktemp("override") / "artifacts"
    run_pipeline("wordpress", root, JOB)
    store = LocalArtifactStore(base_dir=str(root))
    mark_assignment_customer_approved(store, DB, JOB)
    _fake_design(JOB, DB, "dynamodb", store, 2)
    v2 = store.read_json(f"{DB}/{JOB}/assignment/v2/assignment.json")
    flagged = next(
        qa["query_id"] for qa in v2["query_assignments"] if qa["query_id"].startswith(FLAGGED)
    )
    flag = {
        "query_ids": [flagged],
        "flagged_engine": "dynamodb",
        "reason_category": "client_side_join",
        "explanation": "needs a server-side join",
    }
    store.write_json(
        f"{DB}/{JOB}/schema-dynamodb/v2/design_trace.json", {"trace": [], "routing_flags": [flag]}
    )
    orch = LocalOrchestrator(store, llm_mode="bedrock")
    with patch("src.agents.schema_design.handler.run_schema_design_auto", side_effect=_fake_design):
        orch._run_post_schema_routing(JOB, DB)
        apply_assignment_overrides(
            store, DB, JOB, [QueryOverrideInput(query_id=flagged, assigned_engine="dynamodb")]
        )
        for engine in ("dynamodb", "aurora_mysql"):
            _fake_design(JOB, DB, engine, store, 4)
        store.write_json(
            f"{DB}/{JOB}/schema-dynamodb/v4/design_trace.json",
            {"trace": [], "routing_flags": [flag]},
        )
        for engine in ("documentdb", "elasticache", "opensearch"):
            orch._copy_schema_artifacts(JOB, DB, engine, 3, 4)
        orch._run_post_schema_routing(JOB, DB)
    result = PipelineResult(db=DB, job_id=JOB, artifact_root=root)
    common = ["--job-id", JOB, "--db", DB, "--artifact-root", str(root)]
    _run("synthesis", ["scripts/run_synthesis.py", *common, "--llm-mode", "none"], result)
    _run("report", ["scripts/run_report.py", *common], result)
    synthesis = root / DB / JOB / "synthesis" / "v4"
    return {
        "store": store,
        "flagged": flagged,
        "v4": store.read_json(f"{DB}/{JOB}/assignment/v4/assignment.json"),
        "report": json.loads((synthesis / "report.json").read_text()),
        "decision": next(synthesis.glob("*decision-report*.html")).read_text(),
    }


class TestTheCustomersMoveBackHolds:
    def test_a_later_flag_never_reroutes_the_customers_choice(self, overridden):
        store = overridden["store"]
        assert not store.exists(f"{DB}/{JOB}/assignment/v5/assignment.json")
        engine_of = {qa["query_id"]: qa for qa in overridden["v4"]["query_assignments"]}
        assert engine_of[overridden["flagged"]]["assigned_engine"] == "dynamodb"
        assert engine_of[overridden["flagged"]]["customer_override"] is True

    def test_the_undone_move_is_not_listed(self, overridden):
        assert overridden["v4"]["post_schema_reroutes"] == []
        assert overridden["report"]["post_schema_reroutes"] == []
        assert "Moved after your approval" not in overridden["decision"]
