"""Unit tests for the #572 customer-facing wording guard.

Extends the #406 guard (which only checked a bare issue tag survived in a
reality-check reason) to the four leak classes #572 asks for: a raw engine id
where a display name belongs, an ``ADR-`` reference, a snake_case reason code,
and a known log-only phrase. These tests exercise the scanner itself in
isolation, with small, synthetic fixtures; ``tests/e2e/test_customer_facing_wording.py``
runs it against the real wordpress/discourse deterministic pipeline output.
"""

from __future__ import annotations

from src.tools.validation.customer_text_guard import (
    ID_FIELDS,
    KNOWN_REASON_CODES,
    scan_json,
    scan_markdown,
    scan_pptx_text,
    scan_text,
)


class TestEngineIdDetection:
    def test_raw_engine_id_is_flagged(self) -> None:
        findings = scan_text("Workload split: aurora_mysql handles 63 queries.")
        assert [f.kind for f in findings] == ["engine_id"]
        assert findings[0].text == "aurora_mysql"

    def test_display_name_is_not_flagged(self) -> None:
        assert scan_text("Workload split: Aurora MySQL handles 63 queries.") == []

    def test_every_known_engine_id_is_recognised(self) -> None:
        for engine in (
            "dynamodb",
            "documentdb",
            "opensearch",
            "elasticache",
            "neptune",
            "keyspaces",
        ):
            findings = scan_text(f"moved to {engine} today")
            assert [f.text for f in findings] == [engine], engine

    def test_substring_of_a_longer_word_is_not_a_false_positive(self) -> None:
        # "aurora" is a known id, but must not match inside "aurora_mysql" twice,
        # nor inside an unrelated longer word.
        assert scan_text("auroraborealis is unrelated") == []


class TestAdrDetection:
    def test_adr_reference_is_flagged(self) -> None:
        findings = scan_text("Decided per ADR-031.")
        assert [f.kind for f in findings] == ["adr_reference"]
        assert findings[0].text == "ADR-031"

    def test_plain_prose_is_clean(self) -> None:
        assert scan_text("No internal reference here.") == []


class TestSnakeCaseReasonCode:
    def test_reason_code_with_colon_is_flagged(self) -> None:
        findings = scan_text("Status: tie_unresolved: no routing rule decides.")
        assert [f.kind for f in findings] == ["snake_case_reason_code"]
        assert findings[0].text == "tie_unresolved:"

    def test_known_engine_id_followed_by_colon_is_not_double_flagged_as_a_reason_code(
        self,
    ) -> None:
        # "aurora_mysql:" matches the bare snake_case-plus-colon shape too, but it is
        # already reported once as an engine_id finding -- not a second time as a
        # reason code, which would just be noise about the same substring.
        findings = scan_text("Consolidated 3 queries from documentdb to aurora_mysql: ok.")
        kinds = sorted(f.kind for f in findings)
        assert kinds == ["engine_id", "engine_id"]

    def test_ordinary_prose_with_a_colon_is_clean(self) -> None:
        assert scan_text("Note: this is just a sentence.") == []

    def test_a_key_template_is_not_a_reason_code(self) -> None:
        """#572 round 2: a real-model run's ElastiCache key templates
        (wp_options:{option_name}, wp_postmeta:{post_id}:meta,
        wp_woocommerce_order_itemmeta:{order_item_id}:meta) are legitimate
        customer-facing identifiers, not a reason label -- the colon is
        immediately followed by "{", never whitespace."""
        for text in (
            "wp_options:{option_name}",
            "wp_postmeta:{post_id}:meta",
            "wp_woocommerce_order_itemmeta:{order_item_id}:meta",
        ):
            assert scan_text(text) == [], text

    def test_a_reason_code_at_the_end_of_a_sentence_with_no_prose_after_is_not_flagged(
        self,
    ) -> None:
        """ "Only flag a snake_case code used as a reason label ... followed by
        prose" -- a code with nothing (or no letter) after the colon is not
        read as a reason label."""
        assert scan_text("Status: tie_unresolved:") == []
        assert scan_text("Status: tie_unresolved: 42") == []

    def test_a_quoted_identifier_in_backticks_is_not_a_reason_code(self) -> None:
        """#572 round 2: an identifier-shaped token quoted in Markdown backticks
        or an HTML <code> span states "this exact string", not a sentence."""
        assert scan_text("The cache key is `wp_options: option_name` here.") == []
        assert scan_text("The cache key is <code>wp_options: option_name</code> here.") == []

    def test_engine_id_inside_backticks_is_still_flagged(self) -> None:
        """The backtick exemption is scoped to the snake_case detector only --
        a raw engine id is still a leak even quoted."""
        findings = scan_text("Moved to `aurora_mysql` for relational joins.")
        assert [f.kind for f in findings] == ["engine_id"]

    def test_a_planted_reason_code_leak_outside_a_code_span_is_still_caught(self) -> None:
        """A real reason-label leak, planted next to the same key-template shape
        the guard must stay quiet about, is still caught."""
        text = (
            "The cache key `wp_options:{option_name}` was skipped: "
            "tie_unresolved: no routing rule decides between them."
        )
        findings = scan_text(text)
        assert [f.kind for f in findings] == ["snake_case_reason_code"]
        assert findings[0].text == "tie_unresolved:"


class TestSnakeCaseReasonCodeFieldScoping:
    """#587 round 4: a real-model run flagged ordinary prose -- a model-written
    access pattern description mentioning a column name before a colon
    (``"ordered by meta_id: one Query per comment_id"``) -- as a leak. The
    snake_case shape alone can never be stable against free text a model
    writes, so it is now trusted only inside a field this pipeline's own
    deterministic code writes (``REASON_CODE_FIELDS``/``REASON_CODE_FIELD_PATHS``);
    everywhere else the token must also be a ``KNOWN_REASON_CODE``."""

    def test_a_column_name_before_a_colon_in_model_written_prose_is_not_flagged(self) -> None:
        report = {
            "query_groups": [
                {
                    "access_patterns": [
                        {"description": "ordered by meta_id: one Query per comment_id"}
                    ]
                }
            ]
        }
        assert scan_json(report) == []

    def test_a_known_reason_code_is_flagged_even_in_a_model_written_field(self) -> None:
        """Defense in depth: a GENUINE internal code leaking into an
        unexpected field is still caught, even without the deterministic-
        field shape allowance -- it only takes KNOWN_REASON_CODES membership."""
        report = {
            "query_groups": [
                {"access_patterns": [{"description": "tie_unresolved: no routing decided"}]}
            ]
        }
        findings = scan_json(report)
        assert [f.kind for f in findings] == ["snake_case_reason_code"]
        assert findings[0].text == "tie_unresolved:"

    def test_an_unlisted_but_reason_code_shaped_leak_in_a_deterministic_field_is_still_caught(
        self,
    ) -> None:
        """Inside a field this pipeline's own code writes, the shape alone is
        still enough -- a future internal code not yet added anywhere this
        guard imports from is still caught, not silently let through."""
        report = {"assignment_reason": "brand_new_internal_code: not yet in any known list"}
        findings = scan_json(report)
        assert [f.kind for f in findings] == ["snake_case_reason_code"]
        assert findings[0].text == "brand_new_internal_code:"
        assert "brand_new_internal_code" not in KNOWN_REASON_CODES

    def test_cache_notes_list_items_use_the_deterministic_field_shape_allowance(self) -> None:
        report = {"cache_notes": ["5 queries owned by elasticache; tie_unresolved: unresolved"]}
        findings = scan_json(report)
        assert [f.kind for f in findings] == ["engine_id", "snake_case_reason_code"]

    def test_consolidations_and_sweep_decisions_reason_use_the_shape_allowance(self) -> None:
        for report in (
            {"consolidations": [{"reason": "orphan_sweep_internal_code: moved on"}]},
            {"sweep_decisions": [{"reason": "orphan_sweep_internal_code: kept"}]},
        ):
            findings = scan_json(report)
            assert [f.kind for f in findings] == ["snake_case_reason_code"], report

    def test_known_reason_codes_cover_every_imported_source(self) -> None:
        # Spot-check rather than hard-coding the full set, so this test does not
        # have to be rewritten every time a new tie kind/gate kind/blocking rule
        # is added at its own source.
        for code in ("source_compatible", "stability", "tie_unresolved"):
            assert code in KNOWN_REASON_CODES, code
        for code in ("hard_exclusion", "capability"):
            assert code in KNOWN_REASON_CODES, code
        for code in ("signal_override", "no_serviceable_engine", "ineligible_destination", "mixed"):
            assert code in KNOWN_REASON_CODES, code


class TestKnownLogOnlyPhrase:
    def test_exclusion_rule_description_is_flagged(self) -> None:
        # The literal, documented log-only text from one real ExclusionRule
        # (src.agents.referee.engine_exclusions) -- never meant for a customer; its
        # customer_message counterpart is the plain-language version that belongs here.
        findings = scan_text(
            "ElastiCache cannot perform multi-table JOINs; it's a key-value/data " "structure store"
        )
        assert [f.kind for f in findings] == ["log_only_phrase"]

    def test_customer_message_counterpart_is_clean(self) -> None:
        assert scan_text("joins more than one table, and a key-value store can't serve joins") == []


class TestConfidenceWordingDetection:
    """Bug: Deliverables present a fit score as a confidence percentage.

    "NN% confidence" reads as a success probability; the figure is a
    workload-fit score and must be labelled "fit"/"workload fit"
    (``src.shared.workload_fit_labels``). This is a plain regression guard
    against the old wording coming back, not a field exclusion list."""

    def test_percent_confidence_is_flagged(self) -> None:
        findings = scan_text("This database is recommended with 78% confidence based on it.")
        assert [f.kind for f in findings] == ["confidence_wording"]
        assert findings[0].text == "78% confidence"

    def test_percent_fit_is_not_flagged(self) -> None:
        assert scan_text("This database is recommended with 78% fit based on it.") == []

    def test_triage_confidence_label_is_not_flagged(self) -> None:
        # A different metric (triage's classification confidence) this bug does
        # not cover, and its label has no percentage directly before the word.
        assert scan_text("Triage confidence: 87%") == []

    def test_case_insensitive_and_no_space_variant(self) -> None:
        findings = scan_text("Lowest Confidence of the 3 engines: 40%CONFIDENCE here")
        assert [f.text for f in findings] == ["40%CONFIDENCE"]

    def test_confidence_interval_is_not_flagged(self) -> None:
        # Review round 1: a real, unrelated statistics term -- not a mislabelled
        # fit score -- must not trip this guard.
        assert scan_text("This has a 95% confidence interval for the estimate.") == []

    def test_confidence_level_is_not_flagged(self) -> None:
        assert scan_text("This has a 90% confidence level for the test.") == []

    def test_plural_confidence_intervals_is_not_flagged(self) -> None:
        assert scan_text("This has 95% confidence intervals computed per table.") == []

    def test_confidence_interval_case_insensitive(self) -> None:
        assert scan_text("A 95% CONFIDENCE INTERVAL was used.") == []

    def test_a_genuine_leak_right_before_the_word_interval_is_still_caught(self) -> None:
        # The exclusion is scoped to "confidence interval/level" specifically --
        # "confidence" followed by anything else (including a word that merely
        # starts with "interval"-like text) is still flagged.
        findings = scan_text("This database is recommended with 78% confidence, interval unknown.")
        assert [f.kind for f in findings] == ["confidence_wording"]


class TestScanJson:
    def test_id_holding_field_is_excluded_by_path_not_by_value(self) -> None:
        """``assigned_engine`` legitimately holds a raw engine id -- excluded by
        field name, not because the string happens to look like one."""
        assert scan_json({"assigned_engine": "dynamodb"}) == []

    def test_same_raw_id_in_a_prose_field_is_flagged(self) -> None:
        findings = scan_json({"assignment_reason": "highest confidence for dynamodb"})
        assert [f.kind for f in findings] == ["engine_id"]

    def test_dict_keys_are_never_scanned(self) -> None:
        """``owners``-style maps are keyed by raw engine id; only values are prose."""
        assert scan_json({"owners": {"dynamodb": 5, "aurora_mysql": 3}}) == []

    def test_nested_list_of_dicts_is_walked(self) -> None:
        findings = scan_json(
            {"consolidations": [{"from_engine": "documentdb", "reason": "aurora_mysql wins"}]}
        )
        assert [f.kind for f in findings] == ["engine_id"]
        assert findings[0].path == "$.consolidations[0].reason"

    def test_query_text_is_excluded_customers_own_sql_not_internal_wording(self) -> None:
        assert scan_json({"query_text": "SELECT like_count: 1 FROM t"}) == []

    def test_risk_description_s_leading_engine_tag_is_not_a_leak(self) -> None:
        """``Risk.description`` has no separate engine field -- the engine is a
        leading ``[engine]`` tag every renderer parses off before display
        (``src.report.renderers._risk_engine_and_body``). Stripped before
        scanning, not excluded outright, so a leak in the body is still caught."""
        report = {"risks": [{"description": "[aurora_mysql] Table accessed by 2 patterns."}]}
        assert scan_json(report) == []

    def test_a_leak_in_the_risk_description_body_is_still_caught(self) -> None:
        report = {
            "resolved_risks": [{"description": "[aurora_mysql] Consider moving to dynamodb."}]
        }
        findings = scan_json(report)
        assert [f.text for f in findings] == ["dynamodb"]

    def test_tag_stripping_is_qualified_by_parent_not_bare_field_name(self) -> None:
        """#572 review item 5: an unrelated field also named "description"
        (``TradeOff.description``, not a risk) does not use the ``[engine]`` tag
        convention -- a leading "[" there is customer text, not a tag to strip."""
        report = {"trade_offs": [{"description": "[aurora_mysql] is not a tag here."}]}
        findings = scan_json(report)
        assert [f.text for f in findings] == ["aurora_mysql"]

    def test_id_fields_cover_every_contract_engine_id_field(self) -> None:
        # Spot-check a representative sample rather than every field, so this test
        # does not have to be rewritten every time a new id field is added elsewhere.
        for name in (
            "assigned_engine",
            "cache_engine",
            "from_engine",
            "to_engine",
            "recommended_database",
            "primary_engine",
            "tied_engines",
            "engines",
            "read_engines",
            "write_engines",
            "rule_id",
            "signal_override",
            "query_id",
            "query_ids",
            "table_id",
            "source_tables",
            "tables_accessed",
            "query_text",
            "job_id",
            "database_name",
        ):
            assert name in ID_FIELDS, name

    def test_generic_names_are_qualified_by_parent_not_bare(self) -> None:
        """#572 review item 5: these names are common enough that excluding them
        at any depth would hide prose in an unrelated field sharing the name."""
        for name in ("target", "database", "service", "owner", "engine", "names"):
            assert name not in ID_FIELDS, name


class TestScanJsonIdFieldPaths:
    """``ID_FIELD_PATHS`` qualifies a generic field name by its immediate parent
    key (#572 review item 5): the exact (parent, field) pairs below are the real
    ones found in a deterministic wordpress/discourse ``report.json`` (plus
    ``aurora_engine_choice``/``table_owners``, contract fields the samples don't
    happen to populate)."""

    def test_ranking_target_is_excluded(self) -> None:
        assert scan_json({"ranking": [{"target": "dynamodb"}]}) == []

    def test_bare_target_under_an_unrelated_parent_is_still_scanned(self) -> None:
        findings = scan_json({"other_list": [{"target": "dynamodb"}]})
        assert [f.kind for f in findings] == ["engine_id"]

    def test_cost_breakdown_database_is_excluded(self) -> None:
        report = {"tco_analysis": {"cost_breakdown": [{"database": "elasticache"}]}}
        assert scan_json(report) == []

    def test_eliminated_engine_costs_database_is_excluded(self) -> None:
        report = {"tco_analysis": {"eliminated_engine_costs": [{"database": "opensearch"}]}}
        assert scan_json(report) == []

    def test_databases_service_is_excluded(self) -> None:
        report = {"recommended_architecture": {"databases": [{"service": "dynamodb"}]}}
        assert scan_json(report) == []

    def test_bare_database_under_an_unrelated_parent_is_still_scanned(self) -> None:
        findings = scan_json({"some_other_list": [{"database": "dynamodb"}]})
        assert [f.kind for f in findings] == ["engine_id"]

    def test_access_patterns_engine_is_excluded(self) -> None:
        report = {"query_groups": [{"access_patterns": [{"engine": "aurora_mysql"}]}]}
        assert scan_json(report) == []

    def test_cache_overlay_engine_is_excluded(self) -> None:
        assert scan_json({"cache_overlay": {"engine": "elasticache"}}) == []

    def test_aurora_engine_choice_engine_is_excluded(self) -> None:
        assert scan_json({"aurora_engine_choice": {"engine": "aurora_mysql"}}) == []

    def test_resolved_risks_engine_is_excluded(self) -> None:
        assert scan_json({"resolved_risks": [{"engine": "dynamodb"}]}) == []

    def test_trade_offs_engine_is_excluded(self) -> None:
        assert scan_json({"trade_offs": [{"engine": "dynamodb"}]}) == []

    def test_bare_engine_under_an_unrelated_parent_is_still_scanned(self) -> None:
        findings = scan_json({"some_other_list": [{"engine": "dynamodb"}]})
        assert [f.kind for f in findings] == ["engine_id"]

    def test_table_owners_owner_is_excluded(self) -> None:
        assert scan_json({"table_owners": [{"owner": "dynamodb"}]}) == []

    def test_unresolved_names_names_is_excluded(self) -> None:
        assert scan_json({"unresolved_names": {"count": 1, "names": ["dynamodb"]}}) == []

    def test_unresolved_table_names_names_is_excluded(self) -> None:
        report = {"unresolved_table_names": {"count": 1, "names": ["aurora_mysql"]}}
        assert scan_json(report) == []


class TestScanJsonSchemaDesignAndTableMappingIdPaths:
    """#572 round 2: a real-model run (schema design present, which
    ``--llm-mode none`` skips) surfaced ``table_mappings[].alternatives[]``'s
    engine id and target table, and the per-engine ``schema_designs`` projection's
    key/collection/index/data-stream templates, that the deterministic samples
    never exercise. Each case below is the real shape from that run's
    ``report.json`` (see ``tests/unit/report/fixtures/schema_design_wording_fixture.json``,
    trimmed, for the end-to-end version)."""

    def test_table_mapping_alternatives_database_is_excluded(self) -> None:
        report = {
            "table_mappings": [
                {
                    "source_table": "wordpress.wp_actionscheduler_actions",
                    "recommended_database": "dynamodb",
                    "alternatives": [{"database": "elasticache", "target_table": "wp_x:{id}"}],
                    "target_table": "wp_actionscheduler_actions",
                }
            ]
        }
        assert scan_json(report) == []

    def test_schema_design_table_name_key_template_is_excluded(self) -> None:
        report = {
            "schema_designs": {
                "elasticache": {
                    "tables": [{"table_name": "wp_options:{option_name}"}],
                }
            }
        }
        assert scan_json(report) == []

    def test_trade_off_target_tables_is_excluded_at_top_level_and_nested(self) -> None:
        for report in (
            {"trade_offs": [{"target_tables": ["wp_postmeta:{post_id}:meta"]}]},
            {
                "schema_designs": {
                    "elasticache": {
                        "trade_offs": [{"target_tables": ["wp_postmeta:{post_id}:meta"]}]
                    }
                }
            },
        ):
            assert scan_json(report) == [], report

    def test_aurora_optimization_target_and_primary_key_are_excluded(self) -> None:
        report = {
            "schema_designs": {
                "aurora_mysql": {
                    "tables": [{"table_name": "wp_postmeta", "primary_key": ["meta_id"]}],
                    "optimizations": [{"target": "wp_postmeta, wp_woocommerce_order_itemmeta"}],
                }
            }
        }
        assert scan_json(report) == []

    def test_migration_note_object_name_is_excluded(self) -> None:
        report = {
            "schema_designs": {
                "dynamodb": {"migration_notes": [{"object_name": "wp_update_post_trigger"}]}
            }
        }
        assert scan_json(report) == []

    def test_a_genuine_leak_in_an_unrelated_field_under_the_same_parent_is_still_caught(
        self,
    ) -> None:
        """The new exclusions are qualified by parent, like every other
        ID_FIELD_PATHS entry: a real prose field under "tables" (not
        "table_name") is still scanned."""
        report = {
            "schema_designs": {
                "elasticache": {"tables": [{"table_name": "wp_options", "notes": "aurora_mysql"}]}
            }
        }
        findings = scan_json(report)
        assert [f.text for f in findings] == ["aurora_mysql"]


class TestScanMarkdown:
    def test_yaml_frontmatter_keys_are_not_reason_codes(self) -> None:
        text = (
            "---\n"
            'title: "engineering-report"\n'
            'job_id: "e2e-word"\n'
            'source_artifact: "wordpress/e2e-word/synthesis/v2/report.json"\n'
            "---\n"
            "# Report\n"
            "Ordinary body text.\n"
        )
        assert scan_markdown(text) == []

    def test_a_leak_in_the_body_is_still_caught(self) -> None:
        text = "---\n" 'title: "engineering-report"\n' "---\n" "### aurora_mysql\n"
        findings = scan_markdown(text)
        assert [f.kind for f in findings] == ["engine_id"]

    def test_no_frontmatter_is_fine(self) -> None:
        assert scan_markdown("# Report\n\nNo frontmatter here.\n") == []


class TestScanPptxText:
    """#572 review item 4: ``scan_pptx_text`` used to read ``has_text_frame``
    shapes only, which skips a table (``add_table``, no ``text_frame`` of its
    own) and a group shape's members (not direct children of the slide's own
    ``shapes``) entirely."""

    def _deck(self, tmp_path, build) -> str:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        build(slide, Inches)
        path = str(tmp_path / "deck.pptx")
        prs.save(path)
        return path

    def test_table_cell_text_is_scanned(self, tmp_path) -> None:
        def build(slide, Inches):
            table = slide.shapes.add_table(2, 1, Inches(1), Inches(1), Inches(3), Inches(1)).table
            table.cell(0, 0).text = "dynamodb handles 63 queries"
            table.cell(1, 0).text = "Aurora MySQL handles the rest"

        findings = scan_pptx_text(self._deck(tmp_path, build))
        assert [f.text for f in findings] == ["dynamodb"]

    def test_group_shape_members_are_scanned(self, tmp_path) -> None:
        def build(slide, Inches):
            box1 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
            box1.text_frame.text = "clean text"
            box2 = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1))
            box2.text_frame.text = "excluded from elasticache"
            group = slide.shapes.add_group_shape([box1, box2])
            assert group  # the group now owns box1/box2; nothing left ungrouped to re-add

        findings = scan_pptx_text(self._deck(tmp_path, build))
        assert [f.text for f in findings] == ["elasticache"]

    def test_plain_text_frame_still_works(self, tmp_path) -> None:
        def build(slide, Inches):
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
            box.text_frame.text = "ADR-031 decided this"

        findings = scan_pptx_text(self._deck(tmp_path, build))
        assert [f.kind for f in findings] == ["adr_reference"]
