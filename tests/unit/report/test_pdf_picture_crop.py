"""The PDF renderer honours picture crops (issue #240).

``pdf_report`` claims the PDF cannot drift from the deck, but ``draw_picture``
drew the whole image into the frame and ignored ``<a:srcRect>``. A cropped
picture therefore showed regions PowerPoint hides, stretched to fit.
"""

from __future__ import annotations

import io
from typing import Any

from PIL import Image
from pptx.util import Inches

from src.report import pdf_report, pptx_report

RED, GREEN, BLUE, WHITE = (255, 0, 0), (0, 255, 0), (0, 0, 255), (255, 255, 255)


def _quadrants() -> bytes:
    """200x100 PNG: left half red, right half blue, with a green top-right
    quadrant so a vertical crop is observable too."""
    img = Image.new("RGB", (200, 100), RED)
    img.paste(BLUE, (100, 0, 200, 100))
    img.paste(GREEN, (100, 0, 200, 50))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


class _Canvas:
    """Records drawImage calls instead of writing a PDF."""

    def __init__(self) -> None:
        self.calls: list[dict[str, Any]] = []

    def drawImage(self, img, x, y, w, h, **kw) -> None:  # noqa: N802 - reportlab API
        self.calls.append({"img": img, "box": (x, y, w, h)})


def _picture(crop: dict[str, float] | None = None):
    prs = pptx_report.open_deck(keep=1)
    slide = prs.slides[0]
    pic = slide.shapes.add_picture(io.BytesIO(_quadrants()), Inches(1), Inches(1), Inches(2))
    for side, frac in (crop or {}).items():
        setattr(pic, f"crop_{side}", frac)
    return pic


def _drawn(pic) -> tuple[Image.Image, tuple[float, float, float, float]]:
    c = _Canvas()
    pdf_report.draw_picture(c, pic, page_h=540.0)  # type: ignore[arg-type]
    assert len(c.calls) == 1
    reader = c.calls[0]["img"]
    w, h = reader.getSize()
    img = Image.frombytes("RGB", (w, h), reader.getRGBData())
    return img, c.calls[0]["box"]


def test_uncropped_picture_draws_the_whole_image() -> None:
    img, _ = _drawn(_picture())
    assert img.size == (200, 100)


def test_right_crop_keeps_only_the_left_half() -> None:
    pic = _picture({"right": 0.5})
    assert pic._element.blipFill.find(f"{pdf_report.A}srcRect").get("r") == "50000"
    img, box = _drawn(pic)
    assert img.size == (100, 100)
    assert {img.getpixel((0, 0)), img.getpixel((99, 99))} == {RED}
    # The frame itself is unchanged: the crop is drawn into the shape's box.
    assert box[2] == pic.width / pdf_report.EMU_PT


def test_left_and_bottom_crop_keeps_the_top_right_quadrant() -> None:
    img, _ = _drawn(_picture({"left": 0.5, "bottom": 0.5}))
    assert img.size == (100, 50)
    assert {img.getpixel((0, 0)), img.getpixel((99, 49))} == {GREEN}


def test_top_crop() -> None:
    img, _ = _drawn(_picture({"left": 0.5, "top": 0.5}))
    assert img.size == (100, 50)
    assert img.getpixel((50, 25)) == BLUE


def test_negative_crop_is_clamped_to_the_image() -> None:
    # Negative srcRect values pad the image in PowerPoint; the renderer clamps
    # rather than reading outside the bitmap.
    img, _ = _drawn(_picture({"right": -0.25}))
    assert img.size == (200, 100)


def test_cropped_palette_png_with_transparency_renders() -> None:
    # The template's old title logo was an 8-bit palette PNG with a tRNS chunk;
    # reportlab cannot draw a PIL palette image carrying byte transparency, so a
    # crop must hand it a converted image or the whole PDF fails.
    pal = Image.new("RGBA", (200, 100), (255, 0, 0, 0))
    pal.paste((0, 0, 255, 255), (100, 0, 200, 100))
    pal = pal.convert("P")
    pal.info["transparency"] = bytes([0] * 256)
    buf = io.BytesIO()
    pal.save(buf, format="PNG", transparency=bytes([0] * 256))
    prs = pptx_report.open_deck(keep=1)
    pic = prs.slides[0].shapes.add_picture(io.BytesIO(buf.getvalue()), 0, 0, Inches(2))
    pic.crop_left = 0.5
    pdf = pdf_report.pptx_to_pdf(_save(prs))
    assert pdf.startswith(b"%PDF")


def _save(prs) -> bytes:
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
