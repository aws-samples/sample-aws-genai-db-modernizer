"""Neutral title slide, metadata and disclaimer for the executive deck (issue #181).

The deck is generated locally from this public sample, so it must not carry the
AWS Transform wordmark/mark, an "Amazon Confidential" footer, or AWS Transform
metadata. The title slide names the assessment, keeps the AWS logo, and prints the
job ID and generation date so a printed PDF traces back to its run. Every slide's
footer reads "(c) <year> Amazon Web Services, Inc. or its affiliates. For
informational purposes only." with the year the report was generated.
"""

from __future__ import annotations

import datetime as dt
import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from src.report import pdf_report, pptx_report

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "wordpress_report.json"
# Branding, plus Microsoft sensitivity-label properties (docProps/custom.xml): a
# public sample must not carry internal classification metadata.
FORBIDDEN = re.compile(r"transform|confidential|pending_classification|msip_label", re.IGNORECASE)


def _report(timestamp: str | None = "2031-04-05T06:07:08Z") -> dict[str, Any]:
    rep: dict[str, Any] = json.loads(FIXTURE.read_text())
    rep["job_id"] = "job-181"
    if timestamp is None:
        rep.pop("timestamp", None)
    else:
        rep["timestamp"] = timestamp
    return rep


def _footer(year: int) -> str:
    return f"© {year} Amazon Web Services, Inc. or its affiliates. For informational purposes only."


def _forbidden_hits(blob: bytes) -> list[str]:
    hits = []
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                text = z.read(name).decode("utf-8")
                hits += [f"{name}: {m.group(0)}" for m in FORBIDDEN.finditer(text)]
    return hits


def _layout_texts(prs) -> list[str]:
    """Every text frame's text on every slide's layout (where the footer lives)."""
    out = []
    for slide in prs.slides:
        for shape in slide.slide_layout.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return out


@pytest.fixture(scope="module")
def rendered() -> tuple[bytes, bytes]:
    return pdf_report.render_executive_summary_pdf(_report(), {})


class TestTemplate:
    def test_template_parts_carry_no_branding_or_classification_metadata(self) -> None:
        assert _forbidden_hits(pptx_report.TEMPLATE.read_bytes()) == []

    def test_template_author_fields_are_neutral(self) -> None:
        with zipfile.ZipFile(pptx_report.TEMPLATE) as z:
            core = z.read("docProps/core.xml").decode("utf-8")
        allowed = {"", "Database Modernizer Assessment"}
        for tag in ("dc:creator", "cp:lastModifiedBy"):
            m = re.search(rf"<{tag}>([^<]*)</{tag}>|<{tag}/>", core)
            value = (m.group(1) or "") if m else ""
            assert value in allowed, f"{tag}={value!r}"

    def test_template_custom_properties_carry_no_stale_counts(self) -> None:
        # "Slides"/"Notes" described the 43-slide source deck, not this one.
        with zipfile.ZipFile(pptx_report.TEMPLATE) as z:
            custom = z.read("docProps/custom.xml").decode("utf-8")
        assert 'name="Slides"' not in custom
        assert 'name="Notes"' not in custom


class TestRenderedDeck:
    def test_no_branding_or_classification_metadata_in_any_part(self, rendered) -> None:
        deck, _ = rendered
        assert _forbidden_hits(deck) == []

    def test_every_slide_footer_is_the_neutral_disclaimer_with_generation_year(
        self, rendered
    ) -> None:
        prs = Presentation(io.BytesIO(rendered[0]))
        footers = [t for t in _layout_texts(prs) if t.startswith("©")]
        # Each of the seven slides inherits at least one visible footer line.
        assert len(footers) >= len(prs.slides)
        assert set(footers) == {_footer(2031)}

    def test_title_slide(self, rendered) -> None:
        prs = Presentation(io.BytesIO(rendered[0]))
        title = prs.slides[0]
        texts = [s.text_frame.text for s in title.shapes if s.has_text_frame]
        assert "Database Modernizer Assessment" in texts
        small = next(t for t in texts if "job-181" in t)
        assert "2031-04-05" in small
        assert any("wordpress" in t for t in texts)
        # The hexagon mark is gone; nothing on the title slide is custom geometry.
        assert not any(
            s._element.find(f".//{pdf_report.A}custGeom") is not None for s in title.shapes
        )

    def test_title_slide_keeps_the_aws_logo(self, rendered) -> None:
        prs = Presentation(io.BytesIO(rendered[0]))
        pictures = [s for s in prs.slides[0].slide_layout.shapes if s.shape_type == 13]
        # Full-bleed background, the small footer logo, and the AWS logo top-left.
        w, h = int(prs.slide_width or 0), int(prs.slide_height or 0)
        logo = [p for p in pictures if p.top < h // 4 and p.width < w // 4]
        assert len(logo) == 1
        # The plain AWS logo (the same part as the footer logo), uncropped -- not
        # the "aws migrations" lockup cropped down to its "aws" half.
        blip = logo[0]._element.blipFill
        assert logo[0].part.related_part(blip.blip.rEmbed).partname == "/ppt/media/image1.png"
        assert blip.find(f"{pdf_report.A}srcRect") is None
        # Same aspect ratio as the image, so nothing is stretched.
        assert abs(logo[0].width / logo[0].height - 1254 / 750) < 0.01

    def test_unused_wordmark_image_is_not_packaged(self, rendered) -> None:
        with zipfile.ZipFile(io.BytesIO(rendered[0])) as z:
            assert "ppt/media/image3.png" not in z.namelist()

    def test_core_properties_name_the_assessment(self, rendered) -> None:
        cp = Presentation(io.BytesIO(rendered[0])).core_properties
        assert cp.title == "Database Modernizer Assessment — wordpress"
        assert cp.comments == "Database Modernizer Assessment — job job-181 — 2031-04-05T06:07:08Z"
        assert cp.author == "Database Modernizer Assessment"
        assert cp.last_modified_by == "Database Modernizer Assessment"

    def test_each_slide_draws_the_footer_and_footer_logo_once(self, rendered) -> None:
        # The template carried a second footer line + logo hidden behind the
        # full-bleed background on two layouts; the PDF text layer then read the
        # footer twice on those slides.
        prs = Presentation(io.BytesIO(rendered[0]))
        for n, slide in enumerate(prs.slides, start=1):
            art = [s for s in slide.slide_layout.shapes if not s.is_placeholder]
            footers = [s for s in art if s.has_text_frame and s.text_frame.text.startswith("©")]
            assert len(footers) == 1, f"slide {n}"
            small_logos = [
                s for s in art if s.shape_type == 13 and s.top > int(prs.slide_height or 0) * 3 // 4
            ]
            assert len(small_logos) <= 1, f"slide {n}"


def test_footer_year_falls_back_to_the_current_year_without_a_timestamp() -> None:
    deck = pptx_report.render_executive_summary_pptx(_report(timestamp=None), {})
    prs = Presentation(io.BytesIO(deck))
    footers = {t for t in _layout_texts(prs) if t.startswith("©")}
    assert footers == {_footer(dt.datetime.now(dt.UTC).year)}


def test_long_database_name_subtitle_shrinks_to_fit_one_line() -> None:
    rep = _report()
    rep["database_name"] = "customer_orders_and_fulfilment_production_primary"
    prs = Presentation(io.BytesIO(pptx_report.render_executive_summary_pptx(rep, {})))
    sub = next(
        s
        for s in prs.slides[0].shapes
        if s.has_text_frame and "customer_orders" in s.text_frame.text
    )
    run = sub.text_frame.paragraphs[0].runs[0]
    size = run.font.size.pt
    assert size < 32
    width_in = sub.width / 914400
    assert pptx_report._text_em(run.text) * size / 72 <= width_in


@pytest.mark.parametrize("stamp", ["2026-13-45T00:00:00Z", "2026-02-30", "not-a-date"])
def test_invalid_generation_date_is_not_printed(stamp: str) -> None:
    f = pptx_report.derive(_report(timestamp=stamp), {})
    assert f["generated"] == ""
    assert f["year"] == dt.datetime.now(dt.UTC).year


def test_valid_generation_date_is_printed() -> None:
    f = pptx_report.derive(_report(timestamp="2031-04-05T06:07:08Z"), {})
    assert (f["generated"], f["year"]) == ("2031-04-05", 2031)
