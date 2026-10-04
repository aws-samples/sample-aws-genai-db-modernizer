"""The deck uses routed confidence, not the all-analyzed-tables average (#152).

discourse: OpenSearch's analysis average was 2% (311 analyzed tables), so the deck
asked "Confirm OpenSearch? 2% confidence" although the 3 text-search queries routed
to it fit it at 60%. The "Confirm" card, the wave split on ``CONFIDENCE_FLOOR``, the
wave confidence and the ranking rows read ``routed_confidence``; a report written
before it existed keeps ``confidence_score``.
"""

from __future__ import annotations

import io
from typing import Any

from src.report import pptx_report


def _report() -> dict[str, Any]:
    """The discourse shape after #296/#304, with routed confidence."""
    return {
        "database_name": "discourse",
        "job_id": "job-1",
        "timestamp": "2026-10-04T00:00:00Z",
        "ranking": [
            {
                "target": "aurora_postgresql",
                "confidence_score": 67,
                "analysis_confidence": 67,
                "routed_confidence": 93,
                "workload_percent": 76.5,
                "assigned_queries": 1266,
            },
            {
                "target": "dynamodb",
                "confidence_score": 62,
                "analysis_confidence": 62,
                "routed_confidence": 91,
                "workload_percent": 23.3,
                "assigned_queries": 385,
            },
            {
                "target": "opensearch",
                "confidence_score": 2,
                "analysis_confidence": 2,
                "routed_confidence": 60,
                "workload_percent": 0.2,
                "assigned_queries": 3,
            },
            {
                "target": "elasticache",
                "role": "cache_layer",
                "confidence_score": 56,
                "analysis_confidence": 56,
                "routed_confidence": 87,
                "routed_confidence_basis": "cached_reads",
                "workload_percent": 0.0,
                "assigned_queries": 0,
                "cache_overlay_queries": 3,
                "cache_call_share_percent": 25.1,
            },
        ],
        "cache_overlay": {"engine": "elasticache", "query_count": 3, "call_share_percent": 25.1},
        "recommended_architecture": {
            "databases": [
                {"service": "aurora_postgresql", "table_count": 0},
                {"service": "dynamodb", "table_count": 0},
                {"service": "opensearch", "table_count": 0},
                {"service": "elasticache", "table_count": 0},
            ]
        },
        "schema_designs": {},
    }


def _by_target(rep: dict[str, Any], target: str) -> dict[str, Any]:
    return next(r for r in rep["ranking"] if r["target"] == target)


class TestConfirmCard:
    def test_no_confirm_card_when_every_routed_fit_clears_the_floor(self) -> None:
        f = pptx_report.derive(_report(), {})
        assert f["conf"]["opensearch"] == 60
        assert not any(d["question"].startswith("Confirm") for d in f["decisions"])

    def test_confirm_card_shows_the_routed_confidence(self) -> None:
        rep = _report()
        _by_target(rep, "opensearch")["routed_confidence"] = 40
        card = next(
            d
            for d in pptx_report.derive(rep, {})["decisions"]
            if d["question"].startswith("Confirm")
        )
        assert card["question"] == "Confirm OpenSearch?"
        assert card["badge"] == "40% confidence"

    def test_the_analysis_average_no_longer_picks_the_weakest(self) -> None:
        """OpenSearch's 2% analysis average is the lowest, its routed fit is not."""
        rep = _report()
        _by_target(rep, "dynamodb")["routed_confidence"] = 45
        card = next(
            d
            for d in pptx_report.derive(rep, {})["decisions"]
            if d["question"].startswith("Confirm")
        )
        assert card["question"] == "Confirm DynamoDB?"

    def test_legacy_report_keeps_the_analysis_average(self) -> None:
        rep = _report()
        for r in rep["ranking"]:
            r.pop("routed_confidence")
        card = next(
            d
            for d in pptx_report.derive(rep, {})["decisions"]
            if d["question"].startswith("Confirm")
        )
        assert card["badge"] == "2% confidence"


class TestWaves:
    def test_wave_split_reads_the_routed_confidence(self) -> None:
        rep = _report()
        # analysis says under the floor, routed says well above it
        _by_target(rep, "dynamodb")["confidence_score"] = 40
        waves = pptx_report.derive(rep, {})["waves"]
        owner_wave = next(w for w in waves if any(e["engine"] == "dynamodb" for e in w["engines"]))
        assert {e["engine"] for e in owner_wave["engines"]} == {"aurora_postgresql", "dynamodb"}
        assert "confidence from 91%" in owner_wave["note"]

    def test_low_routed_fit_is_sequenced_last(self) -> None:
        rep = _report()
        _by_target(rep, "dynamodb")["routed_confidence"] = 45
        names = [w["names"] for w in pptx_report.derive(rep, {})["waves"]]
        assert names.index("DynamoDB") > names.index("Aurora PostgreSQL")


class TestRuleText:
    def test_rule_says_what_the_confidence_measures(self) -> None:
        f = pptx_report.derive(_report(), {})
        text = pptx_report._sequencing_rule_text(f["engines"], f["conf"], routed=True)
        assert text.startswith(
            "Confidence is the mean fit of the queries routed to each engine "
            "(for the cache layer, of the reads it fronts), not a forecast."
        )

    def test_legacy_rule_text_is_unchanged(self) -> None:
        f = pptx_report.derive(_report(), {})
        text = pptx_report._sequencing_rule_text(f["engines"], f["conf"])
        assert text.startswith("Confidence is the assessment's own measure of evidence strength")


def _signal_only(rep: dict[str, Any]) -> dict[str, Any]:
    """discourse after the #312 review: OpenSearch's 3 queries read ``unknown`` tables."""
    _by_target(rep, "opensearch").update(
        routed_confidence_evidence="signal_only",
        routed_queries=3,
        routed_queries_without_table_evidence=3,
    )
    return rep


class TestSignalOnlyEvidence:
    """#152 decision: the 50% rule stays, but a signal-only score is never solid."""

    def test_signal_only_engine_gets_the_confirm_card_above_the_floor(self) -> None:
        f = pptx_report.derive(_signal_only(_report()), {})
        card = next(d for d in f["decisions"] if d["question"].startswith("Confirm"))
        assert card["question"] == "Confirm OpenSearch?"
        assert card["badge"] == "60% (signal only — no table-level evidence)"
        assert card["against"].startswith("No table-level evidence; confirm the requirement")

    def test_the_weakest_under_the_floor_keeps_the_card_and_names_the_other(self) -> None:
        rep = _signal_only(_report())
        _by_target(rep, "dynamodb")["routed_confidence"] = 45
        cards = [
            d
            for d in pptx_report.derive(rep, {})["decisions"]
            if d["question"].startswith("Confirm")
        ]
        assert [c["question"] for c in cards] == ["Confirm DynamoDB?"]
        assert cards[0]["badge"] == "45% confidence"
        assert cards[0]["against"].endswith("also confirm OpenSearch")

    def test_ranking_rows_and_waves_carry_the_label(self) -> None:
        f = pptx_report.derive(_signal_only(_report()), {})
        assert f["conf_text"]["opensearch"] == "60% (signal only)"
        assert f["conf_text"]["aurora_postgresql"] == "93%"
        search_wave = next(w for w in f["waves"] if w["names"] == "OpenSearch")
        assert search_wave["conf_text"] == "60% (signal only)"
        # the 50% rule is unchanged: 60 is not sequenced as unsure
        assert all("re-scope" not in w["note"] for w in f["waves"])

    def test_partial_evidence_is_labelled_from_a_quarter(self) -> None:
        rep = _report()
        _by_target(rep, "dynamodb").update(
            routed_confidence_evidence="partial",
            routed_queries=385,
            routed_queries_without_table_evidence=100,
        )
        f = pptx_report.derive(rep, {})
        assert f["conf_text"]["dynamodb"] == "91% (partly signal-based)"
        owner_wave = next(w for w in f["waves"] if "DynamoDB" in w["names"])
        assert "confidence from 91%" in owner_wave["note"]  # Aurora (93) is not the minimum
        _by_target(rep, "dynamodb")["routed_queries_without_table_evidence"] = 17
        assert pptx_report.derive(rep, {})["conf_text"]["dynamodb"] == "91%"

    def test_wave_note_labels_a_signal_only_minimum(self) -> None:
        rep = _report()
        _by_target(rep, "dynamodb").update(
            routed_confidence=55, routed_confidence_evidence="signal_only"
        )
        owner_wave = next(
            w for w in pptx_report.derive(rep, {})["waves"] if "DynamoDB" in w["names"]
        )
        assert "confidence from 55% (signal only — no table-level evidence)" in owner_wave["note"]

    def test_the_deck_renders_the_labels(self) -> None:
        from pptx import Presentation  # noqa: PLC0415 - only this test needs it

        data = pptx_report.render_executive_summary_pptx(_signal_only(_report()), {})
        text = " ".join(
            shape.text_frame.text
            for slide in Presentation(io.BytesIO(data)).slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        assert "Confirm OpenSearch?" in text
        assert "60% (signal only — no table-level evidence)" in text


def test_confidence_phrase_keeps_the_label_after_the_word() -> None:
    assert pptx_report._confidence_phrase("93%") == "93% confidence"
    assert pptx_report._confidence_phrase("60% (signal only)") == "60% confidence (signal only)"
