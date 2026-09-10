"""Executive Summary Report — PDF rendering of the PPTX deck.

The customer-facing deliverable in the Artifacts panel is
``summary-executive-report.pdf``; the ``.pptx`` stays on S3 as the editable
source. Both come from one render, because this module does **not** re-lay-out
the slides: it walks the shape tree of the deck ``pptx_report`` just produced
and draws each primitive with reportlab. Positions, sizes, colours, fonts and
wrap widths are read off the actual shapes, so the PDF cannot drift from the
deck — there is one layout implementation, in ``pptx_report``.

Why not LibreOffice: converting through ``soffice`` would add ~600 MB to an
ARM64 agent image for one deliverable, cost a cold start inside the synthesis
path, and stamp non-reproducible PDF dates. reportlab is ~2 MB and, with
``invariant``, byte-deterministic.

What is supported is exactly what the deck contains (verified against it, not
guessed): solid-filled rectangles and rounded rectangles, one gradient-filled
custom-geometry shape (the AWS Transform mark), pictures, tables, and text
frames with per-run size/weight/colour/font, paragraph alignment, space-after
and vertical anchoring. Slide backgrounds resolve ``schemeClr`` with
``lumMod``/``lumOff``. Anything outside that set is skipped rather than
approximated, and logged once at debug level.

Fonts: Amazon Ember Regular and Bold ship in ``assets/fonts``. Amazon Ember
*Display* has no separate file, so headings — always bold in this deck — use
Ember Bold. If a face is unavailable the module falls back to Helvetica so a
PDF is still produced.
"""

from __future__ import annotations

import colorsys
import io
import logging
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Emu
from reportlab.lib.colors import Color
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

from . import pptx_report

logger = logging.getLogger(__name__)

FILENAME = "summary-executive-report.pdf"
FONT_DIR = Path(__file__).parent / "assets" / "fonts"

# EMU per point. python-pptx exposes EMU everywhere; PDF user space is points.
EMU_PT = 12700
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"

# PowerPoint's own default text-frame insets, used when a frame does not set
# them (python-pptx returns None for inherited values).
DEF_MARGIN_LR = Emu(91440)  # 0.1"
DEF_MARGIN_TB = Emu(45720)  # 0.05"

# Line advance as a multiple of point size, and the fraction of that advance
# that sits above the baseline. PowerPoint uses the font's own metrics; Ember's
# ascent/descent put the baseline at ~0.80 of a 1.20 line.
LINE_FACTOR = 1.20
BASELINE_FACTOR = 0.80

_FALLBACK = "Helvetica"
_FALLBACK_BOLD = "Helvetica-Bold"

# Registered reportlab font names, keyed by (typeface, bold).
_FONTS: dict[tuple[str, bool], str] = {}
_FONT_FILES = {
    ("Amazon Ember", False): "AmazonEmber_Rg.ttf",
    ("Amazon Ember", True): "AmazonEmber_Bd.ttf",
    # Amazon Ember Display has no separate file. It is an optical-size variant of
    # the same design, so weight is what has to be preserved: the deck's headings
    # are Display Bold, and the template's 6pt copyright line is Display Regular.
    ("Amazon Ember Display", False): "AmazonEmber_Rg.ttf",
    ("Amazon Ember Display", True): "AmazonEmber_Bd.ttf",
}


def register_fonts() -> None:
    """Register the bundled TTFs once. Missing files degrade to Helvetica."""
    if _FONTS:
        return
    for (face, bold), filename in _FONT_FILES.items():
        path = FONT_DIR / filename
        name = f"Ember-{'Bd' if filename.endswith('Bd.ttf') else 'Rg'}"
        if not path.exists():
            logger.warning("Font %s missing; falling back to %s", path.name, _FALLBACK)
            _FONTS[(face, bold)] = _FALLBACK_BOLD if bold else _FALLBACK
            continue
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))
        _FONTS[(face, bold)] = name


def font_for(typeface: str | None, bold: bool) -> str:
    """Map a PowerPoint typeface + weight to a registered reportlab font.

    Any other Ember variant the template happens to use (Mono, Condensed) maps to
    the base family at the right weight rather than dropping to Helvetica, which
    would change both the look and the measured width mid-deck.
    """
    register_fonts()
    if typeface:
        hit = _FONTS.get((typeface, bold))
        if hit:
            return hit
        if typeface.startswith("Amazon Ember"):
            return _FONTS.get(("Amazon Ember", bold), _FALLBACK_BOLD if bold else _FALLBACK)
    return _FONTS.get(("Amazon Ember", bold)) or (_FALLBACK_BOLD if bold else _FALLBACK)


# ---------------------------------------------------------------------------
# colour
# ---------------------------------------------------------------------------
def rl(rgb: tuple[int, int, int] | None) -> Color | None:
    if rgb is None:
        return None
    return Color(rgb[0] / 255.0, rgb[1] / 255.0, rgb[2] / 255.0)


def lum_transform(rgb: tuple[int, int, int], mod: float, off: float) -> tuple[int, int, int]:
    """Apply OOXML ``lumMod``/``lumOff`` — a luminance scale then offset in HSL.

    The deck's slide background is ``dk1`` (pure black) with lumMod 85% and
    lumOff 15%, i.e. the near-black the whole design sits on. Without this the
    PDF would come out black and the cards would lose their separation.
    """
    h, lightness, s = colorsys.rgb_to_hls(*(c / 255.0 for c in rgb))
    lightness = max(0.0, min(1.0, lightness * mod + off))
    r, g, b = colorsys.hls_to_rgb(h, lightness, s)
    return (round(r * 255), round(g * 255), round(b * 255))


def _theme_colors(prs) -> dict[str, tuple[int, int, int]]:
    """The master's colour scheme, keyed by OOXML name (``dk1``, ``accent1``...).

    The mapping-layer names (``tx1``, ``bg1``, ``tx2``, ``bg2``) are resolved
    through the master's ``clrMap``, not assumed. This deck is a dark theme and
    its map is inverted — ``tx1`` points at ``lt1`` — so hardcoding the usual
    ``tx1``→``dk1`` renders the template's own footer text black on near-black.
    """
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
    master = prs.slide_masters[0]
    xml = master.part.part_related_by(rel).blob.decode("utf-8")
    scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S)
    out: dict[str, tuple[int, int, int]] = {}
    if not scheme:
        return out
    pattern = r"<a:(dk1|dk2|lt1|lt2|accent[1-6]|hlink|folHlink)>(.*?)</a:\1>"
    for tag, body in re.findall(pattern, scheme.group(0), re.S):
        hit = re.search(r'(?:srgbClr val="|lastClr=")([0-9A-Fa-f]{6})', body)
        if hit:
            v = hit.group(1)
            out[tag] = (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))
    clr_map = master._element.find(f"{P}clrMap")
    fallback = {"tx1": "dk1", "tx2": "dk2", "bg1": "lt1", "bg2": "lt2"}
    for alias, default_src in fallback.items():
        src = (clr_map.get(alias) if clr_map is not None else None) or default_src
        if src in out:
            out[alias] = out[src]
    return out


def color_from_element(el, theme: dict[str, tuple[int, int, int]]):
    """Resolve an OOXML colour element (``srgbClr`` or ``schemeClr``) to RGB."""
    if el is None:
        return None
    srgb = el.find(f"{A}srgbClr")
    scheme = el.find(f"{A}schemeClr")
    node = srgb if srgb is not None else scheme
    if node is None:
        return None
    if node is srgb:
        v = node.get("val") or "000000"
        rgb = (int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))
    else:
        rgb = theme.get(node.get("val") or "", (0, 0, 0))
    mod = node.find(f"{A}lumMod")
    off = node.find(f"{A}lumOff")
    if mod is not None or off is not None:
        m = int(mod.get("val") or 100000) / 100000 if mod is not None else 1.0
        o = int(off.get("val") or 0) / 100000 if off is not None else 0.0
        rgb = lum_transform(rgb, m, o)
    return rgb


def run_color(
    run,
    theme: dict[str, tuple[int, int, int]],
    default: tuple[int, int, int] = (255, 255, 255),
) -> tuple[int, int, int]:
    """The run's colour, resolving theme references; ``default`` when inherited.

    The generated slides always set an explicit RGB. The retained intro slide
    is template-authored and uses ``schemeClr`` (LIGHT_2), which is why theme
    resolution exists at all.
    """
    try:
        col = run.font.color
        if col is not None and col.type is not None:
            if col.type == 1:  # RGB
                c = col.rgb
                return (c[0], c[1], c[2])
            name = getattr(col.theme_color, "name", "") or ""
            key = {
                "LIGHT_1": "lt1",
                "LIGHT_2": "lt2",
                "DARK_1": "dk1",
                "DARK_2": "dk2",
                "ACCENT_1": "accent1",
                "ACCENT_2": "accent2",
                "ACCENT_3": "accent3",
                "ACCENT_4": "accent4",
                "ACCENT_5": "accent5",
                "ACCENT_6": "accent6",
                "TEXT_1": "tx1",
                "TEXT_2": "tx2",
                "BACKGROUND_1": "bg1",
                "BACKGROUND_2": "bg2",
            }.get(name)
            if key and key in theme:
                rgb = theme[key]
                brightness = getattr(col, "brightness", 0) or 0
                if brightness:
                    rgb = lum_transform(rgb, 1.0 - abs(brightness), max(brightness, 0.0))
                return rgb
    except (AttributeError, KeyError, TypeError):
        pass
    return default


def shape_fill(shape, theme):
    """Solid fill RGB, or None for anything else (background/gradient/no fill)."""
    spPr = shape._element.find(f".//{A}xfrm/..")
    solid = None
    if spPr is not None:
        solid = spPr.find(f"{A}solidFill")
    if solid is None:
        return None
    return color_from_element(solid, theme)


# ---------------------------------------------------------------------------
# text layout
# ---------------------------------------------------------------------------
class _Tok:
    """A word (or space) carrying the style it must be drawn with."""

    __slots__ = ("text", "font", "size", "color", "width")

    def __init__(self, text: str, font: str, size: float, color: Color, width: float):
        self.text = text
        self.font = font
        self.size = size
        self.color = color
        self.width = width


def _tokens(para, theme, default_size: float, default_color) -> list[_Tok]:
    out: list[_Tok] = []
    for run in para.runs:
        if not run.text:
            continue
        size = run.font.size.pt if run.font.size is not None else default_size
        bold = bool(run.font.bold)
        font = font_for(run.font.name, bold)
        color = rl(run_color(run, theme, default_color)) or Color(1, 1, 1)
        # Split keeping the spaces, so a run boundary mid-phrase does not eat one.
        for piece in re.split(r"(\s+)", run.text):
            if not piece:
                continue
            out.append(_Tok(piece, font, size, color, pdfmetrics.stringWidth(piece, font, size)))
    return out


def _wrap(tokens: list[_Tok], max_w: float) -> list[list[_Tok]]:
    lines: list[list[_Tok]] = []
    cur: list[_Tok] = []
    w = 0.0
    for tok in tokens:
        if tok.text.isspace():
            if cur:
                cur.append(tok)
                w += tok.width
            continue
        if cur and w + tok.width > max_w:
            while cur and cur[-1].text.isspace():
                w -= cur.pop().width
            lines.append(cur)
            cur, w = [tok], tok.width
        else:
            cur.append(tok)
            w += tok.width
    if cur:
        while cur and cur[-1].text.isspace():
            cur.pop()
        lines.append(cur)
    return lines


def _line_width(line: list[_Tok]) -> float:
    return sum(t.width for t in line)


def draw_text_frame(
    c: Canvas,
    tf,
    left: float,
    top: float,
    width: float,
    height: float,
    page_h: float,
    theme,
    *,
    default_size: float = 18.0,
    default_color=(255, 255, 255),
    page_no: int = 1,
) -> None:
    """Lay out a text frame the way PowerPoint does for this deck: wrap inside
    the frame's insets, honour paragraph alignment and space-after, and anchor
    the whole block top or middle."""
    ml = (tf.margin_left if tf.margin_left is not None else DEF_MARGIN_LR) / EMU_PT
    mr = (tf.margin_right if tf.margin_right is not None else DEF_MARGIN_LR) / EMU_PT
    mt = (tf.margin_top if tf.margin_top is not None else DEF_MARGIN_TB) / EMU_PT
    mb = (tf.margin_bottom if tf.margin_bottom is not None else DEF_MARGIN_TB) / EMU_PT
    avail = max(width - ml - mr, 1.0)

    blocks: list[tuple[list[list[_Tok]], Any, float]] = []
    total = 0.0
    for para in tf.paragraphs:
        toks = _tokens(para, theme, default_size, default_color)
        if not toks:
            # Empty paragraph still advances, which is how the deck spaces cards.
            blocks.append(([], para, default_size))
            total += default_size * LINE_FACTOR
            continue
        # A slide-number field renders as the page number, not its literal.
        if len(toks) == 1 and toks[0].text == "\u2039#\u203a":
            toks[0].text = str(page_no)
            toks[0].width = pdfmetrics.stringWidth(toks[0].text, toks[0].font, toks[0].size)
        lines = _wrap(toks, avail)
        size = max(t.size for t in toks)
        blocks.append((lines, para, size))
        total += len(lines) * size * LINE_FACTOR
        total += (para.space_after.pt if para.space_after is not None else 0.0) or 0.0

    anchor = str(tf.vertical_anchor or "")
    box_top = top + mt
    inner_h = max(height - mt - mb, 0.0)
    if "MIDDLE" in anchor and total < inner_h:
        box_top += (inner_h - total) / 2.0

    y = box_top
    for lines, para, size in blocks:
        if not lines:
            y += size * LINE_FACTOR
            continue
        align = str(para.alignment or "")
        for line in lines:
            adv = size * LINE_FACTOR
            baseline = page_h - (y + adv * BASELINE_FACTOR)
            x = left + ml
            if "CENTER" in align:
                x += (avail - _line_width(line)) / 2.0
            elif "RIGHT" in align:
                x += avail - _line_width(line)
            for tok in line:
                if not tok.text.isspace():
                    c.setFillColor(tok.color)
                    c.setFont(tok.font, tok.size)
                    c.drawString(x, baseline, tok.text)
                x += tok.width
            y += adv
        y += (para.space_after.pt if para.space_after is not None else 0.0) or 0.0


# ---------------------------------------------------------------------------
# shapes
# ---------------------------------------------------------------------------
def draw_picture(c: Canvas, shape, page_h: float) -> None:
    try:
        img = ImageReader(io.BytesIO(shape.image.blob))
    except Exception as e:  # noqa: BLE001 — a bad image must not lose the page
        logger.debug("picture %s skipped: %s", shape.name, e)
        return
    x, w = shape.left / EMU_PT, shape.width / EMU_PT
    h = shape.height / EMU_PT
    y = page_h - shape.top / EMU_PT - h
    c.drawImage(img, x, y, w, h, mask="auto")


def draw_custgeom(c: Canvas, shape, page_h: float, theme) -> None:
    """Draw a custom-geometry shape: its path, filled with its gradient.

    One shape in the deck needs this — the AWS Transform mark on the intro
    slide. Dropping it would leave a visible hole where the brand mark sits.
    """
    path_el = shape._element.find(f".//{A}custGeom/{A}pathLst/{A}path")
    if path_el is None:
        return
    pw = float(path_el.get("w") or 0) or 1.0
    ph = float(path_el.get("h") or 0) or 1.0
    x0, w = shape.left / EMU_PT, shape.width / EMU_PT
    h = shape.height / EMU_PT
    y0 = page_h - shape.top / EMU_PT - h

    def pt(node) -> tuple[float, float]:
        px = float(node.get("x") or 0) / pw
        py = float(node.get("y") or 0) / ph
        return (x0 + px * w, y0 + h - py * h)

    p = c.beginPath()
    started = False
    for cmd in path_el:
        tag = cmd.tag.split("}")[1]
        pts = cmd.findall(f"{A}pt")
        if tag == "moveTo" and pts:
            p.moveTo(*pt(pts[0]))
            started = True
        elif tag == "lnTo" and pts and started:
            p.lineTo(*pt(pts[0]))
        elif tag == "cubicBezTo" and len(pts) == 3 and started:
            (c1x, c1y), (c2x, c2y), (ex, ey) = (pt(q) for q in pts)
            p.curveTo(c1x, c1y, c2x, c2y, ex, ey)
        elif tag == "close" and started:
            p.close()
    if not started:
        return

    grad = shape._element.find(f".//{A}gradFill")
    stops: list[tuple[float, tuple[int, int, int]]] = []
    if grad is not None:
        for gs in grad.findall(f"{A}gsLst/{A}gs"):
            rgb = color_from_element(gs, theme)
            if rgb is not None:
                stops.append((int(gs.get("pos") or 0) / 100000, rgb))
    c.saveState()
    c.clipPath(p, stroke=0, fill=0)
    if len(stops) >= 2:
        stops.sort(key=lambda s: s[0])
        ang = 45.0
        lin = grad.find(f"{A}lin") if grad is not None else None
        if lin is not None:
            ang = (int(lin.get("ang") or 0) / 60000.0) % 360.0
        # OOXML measures the gradient vector clockwise from +x with y down;
        # in PDF space y is up, so the vertical component flips.
        import math

        rad = math.radians(ang)
        dx, dy = math.cos(rad), -math.sin(rad)
        cx, cy = x0 + w / 2, y0 + h / 2
        half = max(abs(dx) * w, abs(dy) * h) / 2 or max(w, h) / 2
        c.linearGradient(
            cx - dx * half,
            cy - dy * half,
            cx + dx * half,
            cy + dy * half,
            [rl(s[1]) for s in stops],
            [s[0] for s in stops],
            extend=True,
        )
    else:
        solid = shape_fill(shape, theme)
        if solid:
            c.setFillColor(rl(solid))
            c.rect(x0, y0, w, h, stroke=0, fill=1)
    c.restoreState()


def draw_autoshape(c: Canvas, shape, page_h: float, theme) -> None:
    rgb = shape_fill(shape, theme)
    if rgb is None:
        return
    x, w = shape.left / EMU_PT, shape.width / EMU_PT
    h = shape.height / EMU_PT
    y = page_h - shape.top / EMU_PT - h
    c.setFillColor(rl(rgb))
    name = str(shape.shape_type)
    if "ROUNDED_RECTANGLE" in name:
        try:
            adj = float(shape.adjustments[0])
        except (IndexError, ValueError, TypeError):
            adj = 0.16667
        radius = max(min(adj * min(w, h), min(w, h) / 2), 0.01)
        c.roundRect(x, y, w, h, radius, stroke=0, fill=1)
    else:
        c.rect(x, y, w, h, stroke=0, fill=1)


def draw_table(c: Canvas, shape, page_h: float, theme, page_no: int) -> None:
    tbl = shape.table
    x0 = shape.left / EMU_PT
    y_top = shape.top / EMU_PT
    widths = [col.width / EMU_PT for col in tbl.columns]
    heights = [row.height / EMU_PT for row in tbl.rows]
    for r, row in enumerate(tbl.rows):
        x = x0
        top = y_top + sum(heights[:r])
        for i, cell in enumerate(row.cells):
            w, h = widths[i], heights[r]
            y = page_h - top - h
            rgb = None
            try:
                if cell.fill.type == 1:
                    col = cell.fill.fore_color.rgb
                    rgb = (col[0], col[1], col[2])
            except (AttributeError, TypeError, ValueError):
                rgb = None
            if rgb is not None:
                c.setFillColor(rl(rgb))
                c.rect(x, y, w, h, stroke=0, fill=1)
            draw_text_frame(
                c,
                cell.text_frame,
                x,
                top,
                w,
                h,
                page_h,
                theme,
                default_size=10.0,
                default_color=(241, 243, 243),
                page_no=page_no,
            )
            x += w


def draw_shape(c: Canvas, shape, page_h: float, theme, page_no: int) -> None:
    kind = str(shape.shape_type)
    if shape.has_table:
        draw_table(c, shape, page_h, theme, page_no)
        return
    if "PICTURE" in kind:
        draw_picture(c, shape, page_h)
        return
    if "GROUP" in kind:
        for child in shape.shapes:
            draw_shape(c, child, page_h, theme, page_no)
        return
    if "FREEFORM" in kind:
        draw_custgeom(c, shape, page_h, theme)
        return
    if "AUTO_SHAPE" in kind or "PLACEHOLDER" in kind or "TEXT_BOX" in kind:
        if "TEXT_BOX" not in kind:
            draw_autoshape(c, shape, page_h, theme)
        if shape.has_text_frame and shape.text_frame.text.strip():
            draw_text_frame(
                c,
                shape.text_frame,
                shape.left / EMU_PT,
                shape.top / EMU_PT,
                shape.width / EMU_PT,
                shape.height / EMU_PT,
                page_h,
                theme,
                page_no=page_no,
            )
        return
    logger.debug("shape kind %s not drawn (%s)", kind, shape.name)


def _is_placeholder(shape) -> bool:
    return shape._element.find(f".//{P}ph") is not None


def draw_background(c: Canvas, slide, page_w: float, page_h: float, theme) -> None:
    """Fill the page with the slide's background colour.

    Resolution order is PowerPoint's: the slide's own ``bg``, else its layout's,
    else the master's. The layouts here all carry ``dk1`` + lumMod/lumOff.
    """
    for part in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        bg = part._element.find(f"{P}cSld/{P}bg")
        if bg is None:
            continue
        solid = bg.find(f"{P}bgPr/{A}solidFill")
        rgb = color_from_element(solid, theme) if solid is not None else None
        if rgb is not None:
            c.setFillColor(rl(rgb))
            c.rect(0, 0, page_w, page_h, stroke=0, fill=1)
            return
    c.setFillColor(Color(0.086, 0.114, 0.149))
    c.rect(0, 0, page_w, page_h, stroke=0, fill=1)


def pptx_to_pdf(deck: bytes) -> bytes:
    """Convert the rendered deck to PDF by drawing its own shape tree.

    Deterministic: reportlab is put in invariant mode (no creation timestamp, no
    random document id) and every string drawn comes from the deck, so the same
    ``report.json`` yields byte-identical PDFs.
    """
    register_fonts()
    prs = Presentation(io.BytesIO(deck))
    # slide_width/height are Optional on the part; the template always sets both.
    page_w = int(prs.slide_width or 0) / EMU_PT
    page_h = int(prs.slide_height or 0) / EMU_PT
    theme = _theme_colors(prs)

    buf = io.BytesIO()
    c = Canvas(buf, pagesize=(page_w, page_h), invariant=1, pageCompression=1)
    c.setTitle(prs.core_properties.title or "Executive Summary Report")
    c.setSubject(prs.core_properties.comments or "")
    c.setAuthor("AWS Transform")
    c.setCreator("AWS Transform — Database Modernization Assessment")

    for i, slide in enumerate(prs.slides, start=1):
        draw_background(c, slide, page_w, page_h, theme)
        # Layout art (aurora background, AWS logo, footer rules) draws first.
        # Layout *placeholders* are skipped: their sample copy ("Enter title")
        # only renders through a slide that inherits them, and these slides do
        # not — the deck sets its own title and subtitle shapes.
        for shape in slide.slide_layout.shapes:
            if not _is_placeholder(shape):
                draw_shape(c, shape, page_h, theme, i)
        for shape in slide.shapes:
            draw_shape(c, shape, page_h, theme, i)
        c.showPage()

    c.save()
    data = buf.getvalue()
    logger.info("Rendered %s: %d pages, %d bytes", FILENAME, len(prs.slides._sldIdLst), len(data))
    return data


def render_executive_summary_pdf(
    report: dict[str, Any],
    export_data: dict[str, Any] | None = None,
) -> tuple[bytes, bytes]:
    """Render the deck and its PDF from one pass.

    Returns ``(pptx_bytes, pdf_bytes)`` — the caller keeps the deck as the
    editable source on S3 and publishes the PDF to the Artifacts panel.
    """
    deck = pptx_report.render_executive_summary_pptx(report, export_data)
    return deck, pptx_to_pdf(deck)
